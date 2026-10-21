"""Options sheet for a reviewed DCF: market-implied move, model values vs. the market's lognormal distribution at a long
expiry, chain snapshot, and five strategies with P/L at the model values.  Used by bbuild.py."""
import datetime as dt, json, math, os, re
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as CL

D = os.path.dirname(os.path.abspath(__file__))
SEC = PatternFill("solid", fgColor="F2F2F2")
HDR = PatternFill("solid", fgColor="D9E1F2")
USD = '$#,##0.00;($#,##0.00);"-"'
DOL = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"
RF = 0.04


def edate(code):
    return dt.date(2000 + int(code[:2]), int(code[2:4]), int(code[4:6]))


def add(wb, T, R, put, P):
    raw = json.load(open(f"{D}/rawopt/{T}.json"))
    data = raw["data"]
    quotes = {}
    for o in data["options"]:
        m = re.match(rf"{T}(\d{{6}})([CP])(\d{{8}})$", o["option"])
        if m:
            quotes[(m.group(1), m.group(2), int(m.group(3)) / 1000)] = o
    pdate = dt.date.fromisoformat(data["last_trade_time"][:10])
    spot = float(data["close"])
    q = R.get("q", 0.0)

    def spread(o):
        b, a = o.get("bid") or 0, o.get("ask") or 0
        return (a - b) / ((a + b) / 2) if b > 0 and a > 0 else 9.9

    tight = lambda o: spread(o) <= 0.35          # leg priced at mid
    usable = lambda o: spread(o) <= 0.60 and o.get("iv")     # IV usable
    rv = None
    if os.path.exists(f"{D}/rawopt/{T}_vol.json"):
        rv = json.load(open(f"{D}/rawopt/{T}_vol.json"))["vol6m"]
    IVSRC = {}

    def strikes(e):
        return sorted({k for (ee, cp, k) in quotes if ee == e and (ee, "C", k) in quotes and (ee, "P", k) in quotes})

    def atm_iv(e, k0):
        ivs = [quotes[(e, cp, k0)]["iv"] for cp in "CP" if usable(quotes[(e, cp, k0)])]
        if ivs:
            IVSRC[e] = "ATM quotes"
            return sum(ivs) / len(ivs)
        near = sorted(quotes[(e, cp, k)]["iv"] for k in strikes(e) for cp in "CP"
                      if abs(k / spot - 1) <= 0.15 and usable(quotes[(e, cp, k)]))
        if near:
            IVSRC[e] = "median of near-the-money quotes"
            return near[len(near) // 2]
        IVSRC[e] = "6-month realized volatility (quotes too wide)" if rv else "default 35% (quotes too wide)"
        return rv or 0.35

    def bsm(S, K, Tm, sig, cp):
        d1 = (math.log(S / K) + (RF - q + sig * sig / 2) * Tm) / (sig * math.sqrt(Tm))
        d2 = d1 - sig * math.sqrt(Tm)
        N = lambda x: 0.5 * (1 + math.erf(x / math.sqrt(2)))
        if cp == "C":
            return S * math.exp(-q * Tm) * N(d1) - K * math.exp(-RF * Tm) * N(d2)
        return K * math.exp(-RF * Tm) * N(-d2) - S * math.exp(-q * Tm) * N(-d1)

    en, el = R["near"], R["long"]
    Tn, Tl = (edate(en) - pdate).days / 365, (edate(el) - pdate).days / 365
    kn, kl = strikes(en), strikes(el)
    K0n = min(kn, key=lambda k: (abs(k - spot), -k))
    K0 = min(kl, key=lambda k: (abs(k - spot), -k))
    ivn, ivl = atm_iv(en, K0n), atm_iv(el, K0)

    def prem(e, k, cp, Tm, iv):
        o = quotes[(e, cp, k)]
        if tight(o):
            return (o["bid"] + o["ask"]) / 2, "mid"
        return bsm(spot, k, Tm, iv, cp), "theo"

    bear, base, bull = P["vals"]
    near_k = lambda x: min(kl, key=lambda k: abs(k - x))
    Khi = near_k(base if base > spot * 1.05 else spot * 1.25)
    Khi = Khi if Khi > K0 else min([k for k in kl if k > K0], default=K0)
    Klo = near_k(bear if bear < spot * 0.95 else spot * 0.75)
    Klo = Klo if Klo < K0 else max([k for k in kl if k < K0], default=K0)
    Kp = near_k(spot * 0.90)
    Kp = Kp if Kp < K0 else max([k for k in kl if k < K0], default=K0)

    O = wb.create_sheet("Options")
    O.column_dimensions["A"].width = 44
    for col in "BCDEFGHIJKL":
        O.column_dimensions[col].width = 12
    put(O, "A1", f"{T}: options vs. the reviewed DCF (per contract = 100 shares)", bold=True)
    put(O, "A2", "Market prices are Cboe delayed quotes; blue = input. Probabilities are risk-neutral (lognormal at ATM IV).",
        italic=True)
    put(O, "A3", "Inputs", bold=True, fill=SEC)
    inp = [(4, f"Spot (close {pdate.isoformat()})", spot, USD), (5, "Pricing date", dt.datetime.combine(pdate, dt.time()), DATE),
           (6, "Near expiry", dt.datetime.combine(edate(en), dt.time()), DATE),
           (7, "Long expiry", dt.datetime.combine(edate(el), dt.time()), DATE),
           (8, "Years to near expiry", "=(B6-B5)/365", "0.000"), (9, "Years to long expiry", "=(B7-B5)/365", "0.000"),
           (10, "Risk-free rate", RF, PCT), (11, "Dividend yield", q, PCT),
           (12, "ATM implied volatility, near", round(ivn, 4), PCT), (13, "ATM implied volatility, long", round(ivl, 4), PCT),
           (14, "Contract multiplier", 100, "0")]
    for r, lab, v, fmt in inp:
        put(O, f"A{r}", lab)
        put(O, f"B{r}", v, fmt=fmt)
    # market expectations (near expiry)
    put(O, "A16", "Market expectations, near expiry", bold=True, fill=SEC)
    cn, cnt = prem(en, K0n, "C", Tn, ivn)
    pn, pnt = prem(en, K0n, "P", Tn, ivn)
    rows = [(17, "ATM strike", K0n, "0.00"), (18, f"ATM call ({cnt})", round(cn, 4), USD), (19, f"ATM put ({pnt})", round(pn, 4), USD),
            (20, "Straddle, % of spot (expected move)", "=(B18+B19)/B4", PCT),
            (21, "Expected range, low", "=B4*(1-B20)", USD), (22, "Expected range, high", "=B4*(1+B20)", USD),
            (23, "1-SD move by long expiry (ATM IV)", "=B13*SQRT(B9)", PCT)]
    for r, lab, v, fmt in rows:
        put(O, f"A{r}", lab)
        put(O, f"B{r}", v, fmt=fmt)
    # model vs market at the long expiry
    put(O, "A25", "Model vs. market at the long expiry", bold=True, fill=SEC)
    for i, h in enumerate(["Case", "Value / share", "vs. spot", "Annual return needed", "Market P(above)"]):
        put(O, f"{CL(1 + i)}26", h, bold=True, fill=HDR, align=None if i == 0 else "right")
    put(O, "A27", "Forward price at long expiry")
    put(O, "B27", "=B4*EXP((B10-B11)*B9)", fmt=USD)
    cases = [(28, "Bear", "=Assumptions!D30"), (29, "Base (reviewed)", "=Assumptions!D31"), (30, "Bull", "=Assumptions!D32")]
    for r, lab, f in cases:
        put(O, f"A{r}", lab)
        put(O, f"B{r}", f, fmt=USD)
        put(O, f"C{r}", f"=B{r}/$B$4-1", fmt=PCT)
        put(O, f"D{r}", f'=IF(B{r}>0,(B{r}/$B$4)^(1/$B$9)-1,"")', fmt=PCT)
        put(O, f"E{r}", f"=IF(B{r}>0,NORMSDIST((LN($B$27/B{r})-$B$13^2*$B$9/2)/($B$13*SQRT($B$9))),1)", fmt=PCT)
    # chain snapshot (long expiry)
    put(O, "A32", f"Chain, {edate(el).isoformat()} expiry", bold=True, fill=SEC)
    hd = ["Strike", "Call bid", "Call ask", "Call IV", "Put bid", "Put ask", "Put IV", "Call OI", "Put OI"]
    for i, h in enumerate(hd):
        put(O, f"{CL(1 + i)}33", h, bold=True, fill=HDR, align="center")
    show = sorted(set(sorted(kl, key=lambda k: abs(k - spot))[:8]) | {K0, Khi, Klo, Kp, near_k(base), near_k(bull), near_k(bear)})
    for j, k in enumerate(show):
        r = 34 + j
        c, p_ = quotes[(el, "C", k)], quotes[(el, "P", k)]
        put(O, f"A{r}", k, fmt="0.00", color="000000")
        for col, v, fmt in (("B", c.get("bid") or 0, USD), ("C", c.get("ask") or 0, USD), ("D", c.get("iv") or 0, PCT),
                            ("E", p_.get("bid") or 0, USD), ("F", p_.get("ask") or 0, USD), ("G", p_.get("iv") or 0, PCT),
                            ("H", c.get("open_interest") or 0, "#,##0"), ("I", p_.get("open_interest") or 0, "#,##0")):
            put(O, f"{col}{r}", v, fmt=fmt)
    s0 = 34 + len(show) + 1
    # strategies
    put(O, f"A{s0}", f"Strategies, {edate(el).isoformat()} expiry (premium: + = credit; P/L at expiry per contract)",
        bold=True, fill=SEC)
    legs = {}
    for name, k, cp in (("c0", K0, "C"), ("chi", Khi, "C"), ("p0", K0, "P"), ("plo", Klo, "P"), ("pp", Kp, "P")):
        v, how = prem(el, k, cp, Tl, ivl)
        legs[name] = (k, round(v, 4), how)
    lr = s0 + 1
    put(O, f"A{lr}", "Leg prices", bold=True)
    LG = {}
    for i, (name, lab) in enumerate((("c0", "ATM call"), ("chi", "Upper call"), ("p0", "ATM put"), ("plo", "Lower put"),
                                     ("pp", "Short-put strike"))):
        r = lr + 1 + i
        k, v, how = legs[name]
        put(O, f"A{r}", f"{lab} ({how})")
        put(O, f"B{r}", k, fmt="0.00")
        put(O, f"C{r}", v, fmt=USD)
        LG[name] = (f"$B${r}", f"$C${r}")
    h0 = lr + 7
    heads = ["Strategy", "Premium", "Max profit", "Max loss", "Breakeven", "P/L at bear", "P/L at spot", "P/L at base",
             "P/L at bull", "Market P(profit)"]
    for i, h in enumerate(heads):
        put(O, f"{CL(1 + i)}{h0}", h, bold=True, fill=HDR, align=None if i == 0 else "right")
    K = lambda n: LG[n][0]
    Pm = lambda n: LG[n][1]
    pab = lambda x: f"NORMSDIST((LN($B$27/({x}))-$B$13^2*$B$9/2)/($B$13*SQRT($B$9)))"
    xs = {"F": "$B$28", "G": "$B$4", "H": "$B$29", "I": "$B$30"}
    strat = [
        ("Long ATM call", f"=-{Pm('c0')}*$B$14", '="Unlimited"', "=B{r}", f"={K('c0')}+{Pm('c0')}",
         lambda S: f"=(MAX({S}-{K('c0')},0)-{Pm('c0')})*$B$14", lambda r: f"={pab(f'E{r}')}"),
        ("Bull call spread (ATM / upper)", f"=-({Pm('c0')}-{Pm('chi')})*$B$14", f"=({K('chi')}-{K('c0')})*$B$14+B{{r}}",
         "=B{r}", f"={K('c0')}+{Pm('c0')}-{Pm('chi')}",
         lambda S: f"=(MAX({S}-{K('c0')},0)-MAX({S}-{K('chi')},0)-{Pm('c0')}+{Pm('chi')})*$B$14",
         lambda r: f"={pab(f'E{r}')}"),
        ("Cash-secured put (~10% below spot)", f"={Pm('pp')}*$B$14", "=B{r}", f"=-({K('pp')}*$B$14-B{{r}})",
         f"={K('pp')}-{Pm('pp')}", lambda S: f"=({Pm('pp')}-MAX({K('pp')}-{S},0))*$B$14", lambda r: f"={pab(f'E{r}')}"),
        ("Covered call (100 shares + upper call, incl. dividends)", f"={Pm('chi')}*$B$14",
         f"=({K('chi')}-$B$4)*$B$14+B{{r}}+$B$4*$B$11*$B$9*$B$14", f"=-($B$4*$B$14-B{{r}}-$B$4*$B$11*$B$9*$B$14)",
         f"=$B$4-{Pm('chi')}-$B$4*$B$11*$B$9",
         lambda S: f"=({S}-$B$4+{Pm('chi')}-MAX({S}-{K('chi')},0)+$B$4*$B$11*$B$9)*$B$14", lambda r: f"={pab(f'E{r}')}"),
        ("Bear put spread (ATM / lower)", f"=-({Pm('p0')}-{Pm('plo')})*$B$14", f"=({K('p0')}-{K('plo')})*$B$14+B{{r}}",
         "=B{r}", f"={K('p0')}-{Pm('p0')}+{Pm('plo')}",
         lambda S: f"=(MAX({K('p0')}-{S},0)-MAX({K('plo')}-{S},0)-{Pm('p0')}+{Pm('plo')})*$B$14",
         lambda r: f"=1-{pab(f'E{r}')}")]
    srows = {}
    for i, (lab, pr, mx, ml, be, pl, pp) in enumerate(strat):
        r = h0 + 1 + i
        srows[lab] = r
        put(O, f"A{r}", lab)
        put(O, f"B{r}", pr, fmt=DOL)
        put(O, f"C{r}", mx.replace("{r}", str(r)), fmt=DOL)
        put(O, f"D{r}", ml.replace("{r}", str(r)), fmt=DOL)
        put(O, f"E{r}", be, fmt=USD)
        for col, S in xs.items():
            put(O, f"{col}{r}", pl(S), fmt=DOL)
        put(O, f"J{r}", pp(r), fmt=PCT)
    put(O, f"A{h0 + 7}", "Model output, not a recommendation. Not a licensed advisor.", italic=True)
    return {"spot": spot, "pdate": pdate.isoformat(), "near": edate(en).isoformat(), "long": edate(el).isoformat(),
            "ivn": ivn, "ivl": ivl, "K0n": K0n, "K0": K0, "Khi": Khi, "Klo": Klo, "Kp": Kp, "h0": h0,
            "srows": srows, "legs": {k: v for k, v in legs.items()}, "q": q, "ivsrc": IVSRC}
