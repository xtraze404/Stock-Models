"""Fill NPV / Upside / Model on the Screener's Pharma tab for the basic models and refresh each link's cached Main
values.  usage: bscreener.py SRC OUT"""
import copy, json, os, re, sys
import openpyxl
from openpyxl.workbook.external_link.external import ExternalRow, ExternalCell, ExternalSheetData, ExternalSheetNames

D = os.path.dirname(os.path.abspath(__file__))
SRC, OUT = sys.argv[1], sys.argv[2]
todo = [l.strip() for l in open(f"{D}/todo.txt") if l.strip()]
todo = [t for t in todo if json.load(open(f"{D}/out/{t}_layout.json"))["kind"] != "No data"]   # removed from the Screener
wb = openpyxl.load_workbook(SRC)
ws = wb["Pharma"]
rows = {str(ws.cell(r, 2).value).strip(): r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value}
ref = rows["EVMN"]          # layout v2: J = NPV / Share, K = Upside, L = Model (link index read from G)
n_ok = 0
for t in todo:
    r = rows[t]
    f = ws[f"G{r}"].value
    m = re.search(r"\[(\d+)\]Main!\$C\$4", str(f))
    assert m, (t, f)
    n = int(m.group(1))
    link = wb._external_links[n - 1]
    assert link.file_link.Target.split("/")[-1] == f"{t}.xlsx", (t, link.file_link.Target)
    ws[f"J{r}"] = f'=IF(ISNUMBER([{n}]Main!$C$11),[{n}]Main!$C$11,"")'
    ws[f"K{r}"] = f'=IFERROR(J{r}/F{r}-1,"")'
    ws[f"L{r}"] = f'=[{n}]Main!$C$15&" ("&[{n}]Main!$C$14&")"'
    for col in "JKL":
        s_, d_ = ws[f"{col}{ref}"], ws[f"{col}{r}"]
        d_.font, d_.alignment, d_.border, d_.number_format = (copy.copy(s_.font), copy.copy(s_.alignment),
                                                              copy.copy(s_.border), s_.number_format)
    C = json.load(open(f"{D}/out/{t}_calc.json"))["main"]
    L = json.load(open(f"{D}/out/{t}_layout.json"))
    names = ["Main", "Yearly Model", "Quarterly Model"] + (["Assumptions", "NPV"] if L["kind"] == "DCF" else [])
    cells = []
    for a in ("C3", "C4", "C5", "C7", "C8", "C9", "C11", "C12", "C13", "C14", "C15"):
        v = C.get(a)
        if v is None or v == "":
            continue
        cells.append(ExternalCell(r=a, t="str" if isinstance(v, str) else None, v=str(v)))
    byrow = {}
    for c in cells:
        byrow.setdefault(int(re.sub(r"[A-Z]", "", c.r)), []).append(c)
    book = link.externalBook
    book.sheetNames = ExternalSheetNames(sheetName=names)
    sds = [ExternalSheetData(sheetId=i) for i in range(len(names))]
    sds[0].row = [ExternalRow(r=k, cell=sorted(v, key=lambda c: c.r)) for k, v in sorted(byrow.items())]
    book.sheetDataSet.sheetData = sds
    n_ok += 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("rows filled", n_ok, "saved", OUT)
