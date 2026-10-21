"""Basic NPV workbook from a ticker template: all available fiscal years of history, then
  - DCF  (TTM revenue >= $50M): 10-year forecast, FCF = NOPAT - reinvestment, mid-year discounting, Gordon terminal value,
          Bear/Base/Bull scenarios and a discount x terminal-growth table;
  - Net cash (little or no revenue): value = net cash per share, with burn, runway and market-implied pipeline value;
  - No data: history only, no per-share value.
usage: bbuild.py TICKER  -> out/TICKER.xlsx + out/TICKER_layout.json"""
import copy, datetime as dt, json, math, os, re, shutil, statistics, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.workbook.external_link.external import (ExternalRow, ExternalCell, ExternalSheetData,
                                                      ExternalSheetNames, ExternalSheetDataSet)
from openpyxl.packaging.relationship import Relationship

D = os.path.dirname(os.path.abspath(__file__))
T = sys.argv[1]
PH = "/Users/eric/Documents/Stock Models/Pharma"
SCREENER = "/Users/eric/Documents/Stock Models/Screener.xlsx"
LINKSRC = f"{PH}/LLY.xlsx"
os.makedirs(f"{D}/backup", exist_ok=True)
if not os.path.exists(f"{D}/backup/{T}.xlsx"):
    shutil.copy2(f"{PH}/{T}.xlsx", f"{D}/backup/{T}.xlsx")
SRC, OUT = f"{D}/backup/{T}.xlsx", f"{D}/out/{T}.xlsx"

ROW = {x["t"]: x for x in json.load(open(f"{D}/tickers.json"))}[T]
TPL = json.load(open(f"{D}/templates.json"))[T]
FX = json.load(open(f"{D}/raw/fx.json"))["rates"]
try:
    H = json.load(open(f"{D}/hist/{T}.json"))
except FileNotFoundError:
    H = {"ccy": "USD", "fy": {}, "ttm": {}, "bs_end": None}
# ADS ratios not stated in XBRL (ordinary shares per ADS), from the listing prospectus; 1 = ordinary shares listed
ADS = {"GSK": 2, "TAK": 0.5, "ASND": 1, "TLX": 1, "AAPG": 4, "ADAG": 1.25, "XTLB": 100, "ALPS": 1, "APLM": 1,
       "INCR": 1, "TLSA": 1}
N = 10
REV_MIN = 50.0          # $mm TTM revenue for a DCF
REVIEW = json.load(open(f"{D}/review.json")).get(T) if os.path.exists(f"{D}/review.json") else None

FONT = "Aptos Narrow"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
HDR = PatternFill("solid", fgColor="D9E1F2")
SEC = PatternFill("solid", fgColor="F2F2F2")
TOP = Border(top=Side(style="thin"))
MM1 = '#,##0.0;(#,##0.0);"-"'
MM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
USD = '$#,##0.00;($#,##0.00);"-"'
PLAIN = {"Main", "Yearly Model", "Quarterly Model"}
LINK_RE = re.compile(r"^=('[^']+'|[A-Za-z_]+)!\$?[A-Z]{1,3}\$?\d+$")


def put(ws, addr, value, fmt=None, color=None, bold=False, fill=None, italic=False, align=None, border=None):
    c = ws[addr]
    c.value = value
    if ws.title in PLAIN:
        color = BLACK
    elif color is None:
        if isinstance(value, str) and value.startswith("="):
            color = GREEN if LINK_RE.match(value) else BLACK
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            color = BLUE
        else:
            color = BLACK
    c.font = Font(name=FONT, size=12, bold=bold, italic=italic, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = Alignment(horizontal=align)
    if border:
        c.border = border
    return c


# ------------------------------------------------------------------ inputs
ccy = H["ccy"]
fxr = 1.0 if ccy == "USD" else FX[ccy]          # reporting currency per USD
conv = 1.0                                      # extract.py already converted every fact to USD
price = float(ROW["price"])


def parse_ccy(comment):
    """'Left blank: ... EUR mm: CashAndCashEquivalents 616.041, X 1.0. Source' -> (ccy, [(name, val)])"""
    if not comment:
        return None
    m = re.search(r"\b([A-Z]{3}) mm: (.*?)\.\s+(?:Source|Excludes)", comment)
    if not m:
        return None
    items = [(a, float(b.replace(",", ""))) for a, b in re.findall(r"([A-Za-z]+) (-?[\d,]+\.\d+)", m.group(2))]
    return m.group(1), items


def amount(cell, comment):
    """Template value, or the comment's native-currency figures converted at the Oct 16 rate."""
    v = TPL[cell]
    if v is not None:
        return v, "template"
    p = parse_ccy(comment)
    if p and p[0] in FX and p[1]:
        nums = "+".join(f"{x:g}" for _, x in p[1])
        return f"=({nums})/{FX[p[0]]:g}", f"{p[0]} at {FX[p[0]]:g}/USD"
    if comment and comment.startswith("No standard debt tags"):
        return 0, "no debt tags"
    return None, "n/a"


shares = TPL["C4"]
shares_src = "template"
if shares is None and T in ADS:
    m = re.search(r"SEC count is ([\d,]+) ordinary shares", TPL["C4c"] or "")
    if m:
        ords = float(m.group(1).replace(",", "")) / 1e6
        shares = f"={ords:.6f}/{ADS[T]:g}" if ADS[T] != 1 else round(ords, 6)
        shares_src = f"ADS ratio {ADS[T]:g} (prospectus)"
    else:
        import gzip
        f = json.load(gzip.open(f"{D}/raw/{T}_facts.json.gz"))["facts"]
        xs = f.get("dei", {}).get("EntityCommonStockSharesOutstanding", {}).get("units", {}).get("shares", [])
        if xs:
            e = max(x["end"] for x in xs)
            ords = sum({x.get("accn", i): x["val"] for i, x in enumerate(xs) if x["end"] == e}.values()) / 1e6
            shares = f"={ords:.6f}/{ADS[T]:g}" if ADS[T] != 1 else round(ords, 6)
            shares_src = f"dei {e}; ADS ratio {ADS[T]:g} (prospectus)"
cash, cash_src = amount("C7", TPL["C7c"])
debt, debt_src = amount("C8", TPL["C8c"])


def ev(x):
    """Numeric value of a literal or literal formula like '=1+2' / '=(1+2)/3'."""
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).lstrip("=")
    if not re.fullmatch(r"[\d.+\-*/() eE]+", s):
        raise ValueError(x)
    return float(eval(s))


sh_v, cash_v, debt_v = ev(shares), ev(cash), ev(debt)

# ------------------------------------------------------------------ history ($mm USD)
hist = {}
FIXES = []
for y, r in H["fy"].items():
    g = lambda k: (r.get(k) * conv if r.get(k) is not None else None)
    rev, cogs, rd, sga, oi = g("rev"), g("cogs"), g("rd"), g("sga"), g("oi")
    opex, costs, ni, tax, pretax = g("opex"), g("costs"), g("ni"), g("tax"), g("pretax")
    if oi is None:
        if costs is not None and rev is not None:
            oi = rev - costs
        elif opex is not None:
            base = opex + ((cogs or 0) if opex + 0.5 < (cogs or 0) + (rd or 0) + (sga or 0) else 0)
            oi = (rev or 0) - base
        elif rd is not None or sga is not None:
            oi = (rev or 0) - (cogs or 0) - (rd or 0) - (sga or 0)
    if oi is None and ni is None and rev is None:
        continue
    asset = r.get("assets")
    if rev and asset and rev > 25 * asset and rev / 1000 <= 5 * asset:      # revenue tagged 1000x too large
        FIXES.append(f"FY{y} revenue tagged as ${rev:,.0f}M against ${asset:,.1f}M of total assets; divided by 1,000")
        rev /= 1000
        if cogs and cogs > 25 * asset:
            cogs /= 1000
    rev = rev or 0.0
    if oi is None:                      # only net income: treat it as operating
        oi = ni
    cogs_ = cogs or 0.0
    rd_ = rd or 0.0
    other_op = rev - cogs_ - rd_ - oi   # SG&A & other operating items, ties to reported operating income
    if ni is None:
        ni = oi - (tax or 0)
    hist[int(y)] = dict(rev=rev, cogs=cogs_, rd=rd_, sga=other_op, oi=oi, tax=tax or 0.0, ni=ni,
                        other=ni + (tax or 0.0) - oi, sh=r.get("sh"), end=r["end"], has_cogs=cogs is not None)
years = sorted(hist)
# weighted shares tagged in thousands (or 1000x too large) vs the cover count
if sh_v:
    for y in years:
        s_ = hist[y].get("sh")
        if s_:
            q = sh_v / s_
            if 400 <= q <= 2500:
                hist[y]["sh"] = s_ * 1000
                FIXES.append(f"FY{y} weighted shares tagged in thousands; multiplied by 1,000")
            elif 400 <= 1 / q <= 2500 and "ADS" not in (TPL["C4c"] or ""):
                hist[y]["sh"] = s_ / 1000
                FIXES.append(f"FY{y} weighted shares tagged 1,000x too large; divided by 1,000")
# split-adjust historical share counts (consecutive-year ratio within 2% of an integer >= 2, or its inverse)
for i in range(len(years) - 1, 0, -1):
    a, b = hist[years[i - 1]].get("sh"), hist[years[i]].get("sh")
    if a and b:
        q = b / a
        for k in (q, 1 / q):
            if k >= 1.9 and abs(k - round(k)) / round(k) < 0.02:
                f = round(q) if q > 1 else 1 / round(1 / q)
                for y in years[:i]:
                    if hist[y].get("sh"):
                        hist[y]["sh"] *= f
                break

ttm = {k: v for k, v in H["ttm"].items()}
last = years[-1] if years else None
rev_last = hist[last]["rev"] if last else 0.0


def ttm_val(k):
    t = ttm.get(k)
    if not t or not last or t.get("fy_end") != hist[last]["end"]:
        return None
    return t["val"] * conv


rev_ttm = ttm_val("rev")
if rev_ttm is None:
    rev_ttm = rev_last
oi_ttm = ttm_val("oi")
if oi_ttm is None and ttm_val("rd") is not None and last:
    oi_ttm = rev_ttm - (ttm_val("cogs") or 0) - (ttm_val("rd") or 0) - (ttm_val("sga") or 0)
if oi_ttm is None and last:
    oi_ttm = hist[last]["oi"]
cfo_ttm = ttm_val("cfo")
if cfo_ttm is None and last and H["fy"][str(last)].get("cfo") is not None:
    cfo_ttm = H["fy"][str(last)]["cfo"] * conv
ttm_end = (ttm.get("rev") or {}).get("end") if ttm_val("rev") is not None else (hist[last]["end"] if last else None)

FLAG = None
prev_rev = hist[years[-2]]["rev"] if len(years) >= 2 else 0.0
if rev_ttm and rev_ttm >= REV_MIN and last and (prev_rev < 25 or rev_last > 3 * prev_rev or rev_ttm > 3 * max(rev_last, 1e-9)):
    FLAG = (f"Revenue jumped from ${prev_rev:,.0f}M (FY{years[-2] if len(years) >= 2 else '-'}) to ${rev_last:,.0f}M "
            f"(FY{last}), TTM ${rev_ttm:,.0f}M. That may include one-time licence or milestone income, or a launch "
            f"too recent to extrapolate, so the DCF was not run.")
if sh_v and rev_ttm and rev_ttm >= REV_MIN and last and cash_v is not None and FLAG is None:
    KIND = "DCF"
elif sh_v and cash_v is not None:
    KIND = "Net cash"
else:
    KIND = "No data"

mcap = price * sh_v if sh_v else None
P = {}
if KIND == "DCF":
    # growth: year-to-date YoY if available, else last fiscal-year growth
    t = ttm.get("rev") or {}
    gs, src = [], []
    if t.get("ytd") and t.get("ytd_prior") and t["ytd_prior"] > 0 and t.get("fy_end") == hist[last]["end"]:
        gs.append(t["ytd"] / t["ytd_prior"] - 1); src.append(f"YTD to {t['end']} {gs[-1]:+.0%}")
    for a_, b_ in ((years[-2], last), (years[-3], years[-2])) if len(years) >= 3 else ((years[-2], last),):
        if hist[a_]["rev"] > 0:
            gs.append(hist[b_]["rev"] / hist[a_]["rev"] - 1); src.append(f"FY{b_} {gs[-1]:+.0%}")
    g_obs = min(gs[0], statistics.median(gs)) if gs else 0.05
    g_src = ("lower of the latest and the median of " + ", ".join(src)) if gs else "default"
    g1 = min(max(g_obs, -0.10), 0.35)
    m_ttm = oi_ttm / rev_ttm if rev_ttm else 0.0
    ms = [m_ttm] + [hist[y]["oi"] / hist[y]["rev"] for y in years[-3:] if hist[y]["rev"] > 0]
    m_norm = statistics.median(ms)
    fy_ms = [hist[y]["oi"] / hist[y]["rev"] for y in years[-3:] if hist[y]["rev"] > 0]
    m_fy = statistics.median(fy_ms) if fy_ms else m_ttm
    m0_src = "TTM"
    if m_ttm < m_fy - 0.15:                # TTM hit by one-time charges (impairment, IPR&D)
        m0_raw, m0_src = m_fy, "median of last 3 fiscal years (TTM includes one-time charges)"
    else:
        m0_raw = m_ttm
    m0 = min(max(m0_raw, -1.5), 0.7)
    disc = 0.08 if mcap >= 100000 else 0.085 if mcap >= 20000 else 0.095 if mcap >= 5000 else 0.11 if mcap >= 1000 else 0.125
    PINS = json.load(open(f"{D}/disc_pins.json")) if os.path.exists(f"{D}/disc_pins.json") else {}
    if T in PINS:          # keep the discount rate set at the first build; daily price moves shouldn't change the tier
        disc = PINS[T]
    lc = next((y for y in reversed(years) if hist[y]["has_cogs"] and hist[y]["rev"] > 0), None)
    gm = 1 - hist[lc]["cogs"] / hist[lc]["rev"] if lc else 1.0
    floor = max(min(0.20, gm - 0.15), 0.05)           # margin target floor, lower for low-gross-margin businesses
    m10 = min(max(m_norm, floor), 0.60)
    opx = hist[last]["rd"] + hist[last]["sga"]
    rds = min(max(hist[last]["rd"] / opx, 0.0), 1.0) if opx > 0 else 0.0
    fy_end = dt.date.fromisoformat(hist[last]["end"])
    nxt = fy_end.replace(year=fy_end.year + 1)
    bsd = dt.date.fromisoformat(H["bs_end"]) if H.get("bs_end") else fy_end
    stub = min(max((nxt - bsd).days / 365.0, 0.0), 1.0)
    if stub < 0.05:                      # balance sheet at/after the next FY end: start one year later is not modeled
        stub = 0.05
    P = dict(g_obs=g_obs, g_src=g_src, m_fy=m_fy, m0_src=m0_src, floor=floor, g1=g1, m_ttm=m_ttm, m_norm=m_norm, m0=m0, m10=m10, disc=disc, gm=gm, rds=rds,
             stub=stub, bsd=bsd.isoformat(), tg=0.025, tax=0.17, reinv=0.25,
             scen={"disc": (disc + 0.01, disc, disc - 0.01), "tg": (0.02, 0.025, 0.03),
                   "g1": (g1 - 0.05, g1, g1 + 0.05), "m10": (m10 - 0.05, m10, min(m10 + 0.05, 0.65)),
                   "tax": (0.17, 0.17, 0.17), "reinv": (0.25, 0.25, 0.25)})

if KIND == "DCF" and REVIEW:
    g0 = REVIEW["g"][0]
    P["scen"]["g1"] = (g0 - 0.05, g0, g0 + 0.05)
    P["scen"]["m10"] = tuple(REVIEW["m10"])
    P["g1"], P["m10"] = g0, REVIEW["m10"][1]
    if REVIEW.get("m0") is not None:
        P["m0"], P["m0_src"] = REVIEW["m0"], "reviewed"
    P["review"] = True
# ------------------------------------------------------------------ workbook
wb = openpyxl.load_workbook(SRC)
for _n in list(wb.sheetnames):          # start from the template's three sheets when rebuilding an installed model
    if _n not in ("Main", "Yearly Model", "Quarterly Model"):
        del wb[_n]
main, ym, qm = wb["Main"], wb["Yearly Model"], wb["Quarterly Model"]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.comment is not None:
                c.comment = None

# Yearly Model: clear template columns, history = every year available, then forecast
for row in ym.iter_rows(min_col=2):
    for c in row:
        c.value = None
for rr, lab in {6: "SG&A & Other", 9: "Interest & Other", 18: "Operating Margin", 19: "Actual / Estimate"}.items():
    put(ym, f"A{rr}", lab)
fyears = [last + k for k in range(1, N + 1)] if KIND == "DCF" else []
allyears = years + fyears
YC = {y: CL(2 + i) for i, y in enumerate(allyears)}
for y in allyears:
    put(ym, f"{YC[y]}1", y, bold=True)
    ym.column_dimensions[YC[y]].width = 11
for i in range(len(allyears) + 2, 40):
    ym.column_dimensions[CL(i)].width = 11
for k, y in enumerate(allyears):
    c = YC[y]
    p = YC[allyears[k - 1]] if k > 0 else None
    if y in hist:
        h = hist[y]
        for rr, key in ((2, "rev"), (3, "cogs"), (5, "rd"), (6, "sga"), (9, "other"), (11, "tax"), (12, "ni")):
            put(ym, f"{c}{rr}", round(h[key], 3), fmt=MM1)
        put(ym, f"{c}14", round(h["sh"], 3) if h.get("sh") else None, fmt=MM1)
        put(ym, f"{c}19", "A", align="right")
        put(ym, f"{c}12", round(h["ni"], 3), fmt=MM1)
    else:
        nc = CL(3 + fyears.index(y))          # NPV column for this forecast year
        put(ym, f"{c}2", f"={p}2*(1+NPV!{nc}5)", fmt=MM1)
        put(ym, f"{c}3", f"={c}2*(1-Assumptions!$B$19)", fmt=MM1)
        put(ym, f"{c}5", f"={c}7*Assumptions!$B$20", fmt=MM1)
        put(ym, f"{c}6", f"={c}7-{c}5", fmt=MM1)
        put(ym, f"{c}9", 0, fmt=MM1)
        put(ym, f"{c}11", f"=MAX(0,{c}10)*Assumptions!$E$14", fmt=MM1)
        put(ym, f"{c}12", f"={c}10-{c}11", fmt=MM1)
        put(ym, f"{c}14", "=Main!$C$4", fmt=MM1)
        put(ym, f"{c}19", "E", align="right")
    put(ym, f"{c}4", f"={c}2-{c}3", fmt=MM1)
    if y in hist:
        put(ym, f"{c}7", f"=SUM({c}5:{c}6)", fmt=MM1)
    else:
        put(ym, f"{c}7", f"={c}4-{c}2*NPV!{CL(3 + fyears.index(y))}6", fmt=MM1)
    put(ym, f"{c}8", f"={c}4-{c}7", fmt=MM1)
    put(ym, f"{c}10", f"={c}8+{c}9", fmt=MM1)
    put(ym, f"{c}13", f"=IFERROR({c}12/{c}14,0)", fmt="0.00")
    if p:
        put(ym, f"{c}16", f'=IFERROR({c}2/{p}2-1,"")', fmt="0.0%")
    put(ym, f"{c}17", f'=IFERROR({c}4/{c}2,"")', fmt="0.0%")
    put(ym, f"{c}18", f'=IFERROR({c}8/{c}2,"")', fmt="0.0%")
ym.column_dimensions["A"].width = 22
if not allyears:
    put(ym, "B19", "No annual financials in SEC XBRL")

sheets = [main, ym, qm]
L = {"kind": KIND, "years": years, "fyears": fyears, "ccy": ccy, "fxr": fxr, "shares_src": shares_src,
     "cash_src": cash_src, "debt_src": debt_src, "rev_ttm": rev_ttm, "oi_ttm": oi_ttm, "cfo_ttm": cfo_ttm,
     "ttm_end": ttm_end, "bs_end": H.get("bs_end"), "flag": FLAG, "fixes": FIXES, "mcap": mcap, "sh": sh_v, "cash": cash_v, "debt": debt_v,
     "price": price, "P": {k: v for k, v in P.items()},
     "hist": {y: {k: v for k, v in h.items() if k != "has_cogs"} for y, h in hist.items()}}

if KIND == "DCF":
    A = wb.create_sheet("Assumptions")
    NP = wb.create_sheet("NPV")
    sheets += [A, NP]
    A.column_dimensions["A"].width = 44
    for col in "BCDEF":
        A.column_dimensions[col].width = 14
    put(A, "A1", f"{ROW['name']}: basic DCF assumptions", bold=True)
    put(A, "A2", "Blue = input, black = formula, green = link. $mm unless noted.", italic=True)
    put(A, "A4", "Scenario (1 Bear, 2 Base, 3 Bull)", bold=True)
    put(A, "B4", 2, fmt="0")
    put(A, "C4", '=CHOOSE($B$4,"Bear","Base","Bull")', bold=True)
    put(A, "A6", "Discount rate override (blank = scenario)")
    put(A, "B6", None, fmt=PCT, color=BLUE)
    put(A, "A7", "Terminal growth override (blank = scenario)")
    put(A, "B7", None, fmt=PCT, color=BLUE)
    for i, h in enumerate(["Parameter", "Bear", "Base", "Bull", "Live"]):
        put(A, f"{CL(1 + i)}9", h, bold=True, fill=HDR, align=None if i == 0 else "right")
    rows = [(10, "Discount rate", "disc", PCT), (11, "Terminal growth", "tg", PCT),
            (12, "Revenue growth, first forecast year", "g1", PCT), (13, f"Operating margin, year {N}", "m10", PCT),
            (14, "Tax rate on operating income", "tax", PCT), (15, "Reinvestment, % of revenue increase", "reinv", PCT)]
    for rr, lab, key, fmt in rows:
        put(A, f"A{rr}", lab)
        for j, v in enumerate(P["scen"][key]):
            put(A, f"{CL(2 + j)}{rr}", round(v, 4), fmt=fmt)
        live = f"CHOOSE($B$4,B{rr},C{rr},D{rr})"
        if key == "disc":
            live = f"IF(ISNUMBER($B$6),$B$6,{live})"
        if key == "tg":
            live = f"IF(ISNUMBER($B$7),$B$7,{live})"
        put(A, f"E{rr}", f"={live}", fmt=fmt, bold=True)
    put(A, "A17", "Single inputs", bold=True, fill=SEC)
    for col in "BCDE":
        A[f"{col}17"].fill = SEC
    singles = [(18, "Starting operating margin (TTM)", round(P["m0"], 4), PCT),
               (19, "Gross margin (forecast)", round(P["gm"], 4), PCT),
               (20, "R&D share of operating expense (forecast)", round(P["rds"], 4), PCT),
               (21, "Forecast years", N, "0"),
               (22, "First forecast year, fraction after balance sheet date", round(P["stub"], 4), "0.00"),
               (23, "Balance sheet date", dt.datetime.fromisoformat(P["bsd"]), "yyyy-mm-dd"),
               (24, f"Reporting currency per USD ({ccy})", fxr, "0.0000"),
               (25, f"Base revenue, FY{last}", f"='Yearly Model'!{YC[last]}2", MM1)]
    for rr, lab, v, fmt in singles:
        put(A, f"A{rr}", lab)
        put(A, f"B{rr}", v, fmt=fmt, color=BLACK if rr in (21, 24) else None)
    if REVIEW:
        put(A, "A44", "Reviewed revenue growth by year (base case; blank = fade formula)", bold=True, fill=SEC)
        for k, y in enumerate(fyears):
            put(A, f"{CL(2 + k)}45", y, fmt="0", bold=True, fill=HDR, align="right")
            put(A, f"{CL(2 + k)}46", REVIEW["g"][k], fmt=PCT)
        put(A, "A45", "Fiscal year")
        put(A, "A46", "Growth (bear/bull shift by the row-12 gap, fading to 0 by year 10)")
    # NPV sheet
    NP.column_dimensions["A"].width = 38
    NP.column_dimensions["B"].width = 13
    put(NP, "A1", f"{ROW['name']}: free-cash-flow DCF ($mm)", bold=True)
    labels = {3: "Fiscal year", 4: "Period", 5: "Revenue growth", 6: "Operating margin", 7: "Revenue",
              8: "Operating income", 9: "Tax on operating income", 10: "NOPAT", 11: "Reinvestment",
              12: "Free cash flow", 13: "Year fraction", 14: "Discount time (mid-period)", 15: "Discount factor",
              16: "PV of free cash flow"}
    for rr, lab in labels.items():
        put(NP, f"A{rr}", lab, bold=rr in (3, 12, 16))
    put(NP, "B3", f"='Yearly Model'!{YC[last]}1", fmt="0", bold=True, fill=HDR)
    put(NP, "B4", "Actual", align="right")
    put(NP, "B7", f"='Yearly Model'!{YC[last]}2", fmt=MM1)
    put(NP, "B8", f"='Yearly Model'!{YC[last]}8", fmt=MM1)
    put(NP, "B6", f"=IFERROR(B8/B7,0)", fmt=PCT)
    for k, y in enumerate(fyears, start=1):
        c, p = CL(2 + k), CL(1 + k)
        NP.column_dimensions[c].width = 11
        put(NP, f"{c}3", f"='Yearly Model'!{YC[y]}1", fmt="0", bold=True, fill=HDR)
        put(NP, f"{c}4", k, fmt="0", color=BLACK)
        fade = f"Assumptions!$E$12+(Assumptions!$E$11-Assumptions!$E$12)*({c}4-1)/(Assumptions!$B$21-1)"
        if REVIEW:
            ov = f"Assumptions!{CL(1 + k)}$46"
            fade = (f"IF(ISNUMBER({ov}),{ov}+(Assumptions!$E$12-Assumptions!$C$12)*(Assumptions!$B$21-{c}4)/"
                    f"(Assumptions!$B$21-1),{fade})")
        put(NP, f"{c}5", f"={fade}", fmt=PCT)
        put(NP, f"{c}6", f"=Assumptions!$B$18+(Assumptions!$E$13-Assumptions!$B$18)*{c}4/Assumptions!$B$21", fmt=PCT)
        put(NP, f"{c}7", f"='Yearly Model'!{YC[y]}2", fmt=MM1)
        put(NP, f"{c}8", f"='Yearly Model'!{YC[y]}8", fmt=MM1)
        put(NP, f"{c}9", f"=MAX(0,{c}8)*Assumptions!$E$14", fmt=MM1)
        put(NP, f"{c}10", f"={c}8-{c}9", fmt=MM1)
        put(NP, f"{c}11", f"=({c}7-{p}7)*Assumptions!$E$15", fmt=MM1)
        put(NP, f"{c}12", f"={c}10-{c}11", fmt=MM1, bold=True)
        put(NP, f"{c}13", "=Assumptions!$B$22" if k == 1 else 1, fmt="0.00", color=None if k == 1 else BLACK)
        put(NP, f"{c}14", f"={c}13/2" if k == 1 else f"={p}14+{p}13/2+{c}13/2", fmt="0.00")
        put(NP, f"{c}15", f"=1/(1+Assumptions!$E$10)^{c}14", fmt="0.000")
        put(NP, f"{c}16", f"={c}12*{c}13*{c}15", fmt=MM1, bold=True)
    LC = CL(2 + N)
    summ = [(19, "PV of forecast free cash flow", f"=SUM(C16:{LC}16)", MM1),
            (20, f"Terminal free cash flow (year {N + 1})",
             f"={LC}10*(1+Assumptions!$E$11)-{LC}7*Assumptions!$E$11*Assumptions!$E$15", MM1),
            (21, "Terminal value", "=B20/(Assumptions!$E$10-Assumptions!$E$11)", MM1),
            (22, "PV of terminal value", f"=B21/(1+Assumptions!$E$10)^({LC}14+0.5)", MM1),
            (23, "Enterprise value", "=B19+B22", MM1),
            (24, "Cash & investments", "=Main!C7", MM1),
            (25, "Debt", "=Main!C8", MM1),
            (26, "Equity value", "=B23+B24-B25", MM1),
            (27, "Shares (mm)", "=Main!C4", MM1),
            (28, "NPV / share", "=B26/B27", USD),
            (29, "Price", "=Main!C3", USD),
            (30, "Upside", "=B28/B29-1", PCT),
            (31, "Terminal value, % of EV", "=IFERROR(B22/B23,0)", PCT)]
    for rr, lab, f, fmt in summ:
        put(NP, f"A{rr}", lab, bold=rr in (23, 26, 28))
        put(NP, f"B{rr}", f, fmt=fmt, bold=rr in (23, 26, 28), border=TOP if rr in (23, 26, 28) else None)
    # scenario summary (data table on B4) and discount x terminal growth (B6, B7)
    put(A, "A27", "Scenario summary", bold=True, fill=SEC)
    for col in "BCDE":
        A[f"{col}27"].fill = SEC
    for i, (h, f, fmt) in enumerate([("Enterprise value", "=NPV!B23", MM), ("Equity value", "=NPV!B26", MM),
                                     ("NPV / share", "=NPV!B28", USD), ("Upside", "=NPV!B30", PCT)]):
        put(A, f"{CL(2 + i)}28", h, bold=True, fill=HDR, align="right")
        put(A, f"{CL(2 + i)}29", f, fmt=fmt, italic=True)
        for j in range(3):
            A[f"{CL(2 + i)}{30 + j}"].number_format = fmt
            A[f"{CL(2 + i)}{30 + j}"].font = Font(name=FONT, size=12)
    put(A, "A28", "Scenario", bold=True, fill=HDR)
    put(A, "A29", "Live", italic=True)
    for j in range(3):
        put(A, f"A{30 + j}", j + 1, fmt='[=1]"Bear";[=2]"Base";"Bull"', color=BLACK)
    A["B30"] = DataTableFormula(ref="B30:E32", r1="B4", dt2D=False, dtr=False)
    put(A, "A35", "NPV / share: discount rate (rows) x terminal growth (columns)", bold=True, fill=SEC)
    for col in "BCDEF":
        A[f"{col}35"].fill = SEC
    put(A, "A36", "=NPV!B28", fmt=USD, italic=True)
    tgs = [0.015, 0.02, 0.025, 0.03, 0.035]
    drs = [round(disc_ + d, 4) for disc_ in [P["disc"]] for d in (-0.02, -0.01, 0, 0.01, 0.02)]
    for i, v in enumerate(tgs):
        put(A, f"{CL(2 + i)}36", v, fmt=PCT, bold=True, fill=HDR)
    for j, v in enumerate(drs):
        put(A, f"A{37 + j}", v, fmt=PCT, bold=True, fill=HDR, align="right")
        for i in range(5):
            A[f"{CL(2 + i)}{37 + j}"].number_format = USD
            A[f"{CL(2 + i)}{37 + j}"].font = Font(name=FONT, size=12)
    A["B37"] = DataTableFormula(ref="B37:F41", r1="B7", r2="B6", dt2D=True)
    L.update(tgs=tgs, drs=drs)

# ------------------------------------------------------------------ Main
for rr in range(10, 25):
    for col in "BC":
        main[f"{col}{rr}"].value = None
put(main, "C3", f'=INDEX([1]Pharma!$F$1:$F$1000,MATCH("{T}",[1]Pharma!$B$1:$B$1000,0))', fmt="0.00")
put(main, "C4", shares, fmt="#,##0.0")
put(main, "C7", cash, fmt="#,##0.0")
put(main, "C8", debt, fmt="#,##0.0")
put(main, "C5", "=C3*C4", fmt="#,##0")
put(main, "C9", "=C5-C7+C8", fmt="#,##0")
lab11 = {"DCF": "NPV / Share", "Net cash": "Net Cash / Share", "No data": "NPV / Share"}[KIND]
items = [(11, lab11, {"DCF": "=NPV!B28", "Net cash": "=(C7-C8)/C4", "No data": None}[KIND], "0.00"),
         (12, "Upside", '=IFERROR(C11/C3-1,"")' if KIND != "No data" else None, "0.0%"),
         (13, "Cash / Share", '=IFERROR((C7-C8)/C4,"")', "0.00"),
         (14, "Scenario", {"DCF": "=Assumptions!C4", "No data": "no data"}.get(KIND, "revenue jump" if FLAG else
                                                                     "pre-revenue" if (rev_ttm or 0) < 5 else "small revenue"), "General"),
         (15, "Model", {"DCF": "Reviewed DCF" if REVIEW else "DCF", "Net cash": "Net cash", "No data": "n/a"}[KIND], "General")]
if KIND != "DCF":
    items += [(17, "TTM Operating Cash Flow", round(cfo_ttm, 3) if cfo_ttm is not None else None, "#,##0.0"),
              (18, "Runway (years)", '=IF(AND(ISNUMBER(C17),C17<0),IF(C7>C8,(C7-C8)/-C17,"net debt"),"n/a")', "0.0"),
              (19, "Implied Pipeline Value", "=C9", "#,##0")]
for rr, lab, f, fmt in items:
    put(main, f"B{rr}", lab)
    put(main, f"C{rr}", f, fmt=fmt)

# external link to Screener.xlsx (cache: tickers + prices from the Screener on disk)
link = copy.deepcopy(openpyxl.load_workbook(LINKSRC)._external_links[0])
link.file_link = Relationship(type="externalLinkPath", Target="../Screener.xlsx", TargetMode="External", Id="rId1")
link.externalBook.id = "rId1"
SC = json.load(open(f"{D}/screener_cache.json"))
rows_ = []
for i, tick, pr in SC["rows"]:
    cells = [ExternalCell(r=f"B{i}", t="str", v=str(tick))]
    if pr is not None:
        cells.append(ExternalCell(r=f"F{i}", t="str" if isinstance(pr, str) else None, v=str(pr)))
    rows_.append(ExternalRow(r=i, cell=cells))
names_ = SC["sheets"]
link.externalBook.sheetNames = ExternalSheetNames(sheetName=names_)
sds = [ExternalSheetData(sheetId=i) for i in range(len(names_))]
sds[names_.index("Pharma")].row = rows_
link.externalBook.sheetDataSet = ExternalSheetDataSet(sheetData=sds)
wb._external_links = [link]

if KIND == "DCF" and REVIEW and os.path.exists(f"{D}/out/{T}_calc.json"):
    import bopts
    S_ = json.load(open(f"{D}/out/{T}_calc.json")).get("scen")
    if S_:
        P["vals"] = [S_[k]["ps"] for k in ("Bear", "Base", "Bull")]
        L["opt"] = bopts.add(wb, T, REVIEW, put, P)
        sheets.append(wb["Options"])
wb._sheets = sheets
wb.active = 0
wb.calculation.fullCalcOnLoad = True
os.makedirs(f"{D}/out", exist_ok=True)
wb.save(OUT)
json.dump(L, open(f"{D}/out/{T}_layout.json", "w"), indent=0, default=str)
print(f"{T}: {KIND} | years {years[:1]}..{years[-1:]} | rev TTM {rev_ttm:,.0f} | sh {sh_v} cash {cash_v} debt {debt_v}"
      + (f" | g1 {P['g1']:.1%} m0 {P['m0']:.1%} m10 {P['m10']:.1%} r {P['disc']:.1%} stub {P['stub']:.2f}" if P else ""))
