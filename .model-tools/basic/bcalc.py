"""Recalculate a basic NPV workbook with the `formulas` engine (Excel not used), compute its data tables and the
price-implied year-10 margin, write cached values into every formula cell, fix the external-link rels Id.
usage: bcalc.py TICKER  -> cached/TICKER.xlsx + out/TICKER_calc.json"""
import json, os, re, sys, time, zipfile
from xml.sax.saxutils import escape
import openpyxl, formulas
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl

D = os.path.dirname(os.path.abspath(__file__))
T = sys.argv[1]
L = json.load(open(f"{D}/out/{T}_layout.json"))
src, vpath, dst = f"{D}/out/{T}.xlsx", f"{D}/calc/{T.lower()}_v.xlsx", f"{D}/cached/{T}.xlsx"
os.makedirs(f"{D}/calc", exist_ok=True)
os.makedirs(f"{D}/cached", exist_ok=True)
t0 = time.time()
wb = openpyxl.load_workbook(src)
wb["Main"]["C3"].value = L["price"]
wb._external_links = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, DataTableFormula):
                c.value = None
wb.save(vpath)
fname = os.path.basename(vpath)
X = lambda s, a: f"'[{fname}]{s.upper()}'!{a}".replace("$", "")
xl = formulas.ExcelModel().loads(vpath).finish()


def val(v):
    x = v.value[0, 0] if hasattr(v, "value") else v
    x = x.item() if hasattr(x, "item") else x
    if x.__class__.__name__ in ("XlError",) or (isinstance(x, str) and x.startswith("#")):
        return str(x)
    return x


sol = xl.calculate()
vals = {k: val(v) for k, v in sol.items()}
errs = [(k, v) for k, v in vals.items() if isinstance(v, str) and v.startswith("#")]
OUT = {"errors": errs[:20], "n_cells": len(vals)}
g = lambda s, a: vals.get(X(s, a))
OUT["main"] = {a: g("Main", a) for a in ("C3", "C4", "C5", "C7", "C8", "C9", "C11", "C12", "C13", "C14", "C15",
                                         "C17", "C18", "C19")}

zin = zipfile.ZipFile(src)
wbx = zin.read("xl/workbook.xml").decode()
rels = zin.read("xl/_rels/workbook.xml.rels").decode()
rid = {m.group(1): m.group(2) for m in re.finditer(r'Id="(rId\d+)"[^>]*?Target="([^"]+)"', rels)}
rid.update({m.group(2): m.group(1) for m in re.finditer(r'Target="([^"]+)"[^>]*?Id="(rId\d+)"', rels)})
sheets = {}
for name, r in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx):
    t = rid[r].lstrip("/")
    sheets["xl/" + t if not t.startswith("xl/") else t] = name.replace("&amp;", "&")

if L["kind"] == "DCF":
    run = lambda inp, out: val(xl.calculate(inputs=inp, outputs=[out])[out])
    # data tables, computed the way Excel does
    for path, sname in sheets.items():
        xml = zin.read(path).decode()
        for m in re.finditer(r'<f t="dataTable"([^>]*)/>', xml):
            a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
            (c0, r0), (c1, r1) = [(ci(re.match(r"[A-Z]+", x).group()), int(re.search(r"\d+", x).group()))
                                  for x in a["ref"].split(":")]
            hv = lambda c, r: vals.get(X(sname, f"{cl(c)}{r}"))
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    if a.get("dt2D") == "1":
                        inp, out = {X(sname, a["r1"]): hv(c, r0 - 1), X(sname, a["r2"]): hv(c0 - 1, r)}, \
                            X(sname, f"{cl(c0 - 1)}{r0 - 1}")
                    elif a.get("dtr") == "1":
                        inp, out = {X(sname, a["r1"]): hv(c, r0 - 1)}, X(sname, f"{cl(c0 - 1)}{r}")
                    else:
                        inp, out = {X(sname, a["r1"]): hv(c0 - 1, r)}, X(sname, f"{cl(c)}{r0 - 1}")
                    vals[X(sname, f"{cl(c)}{r}")] = run(inp, out)
    if "opt" in L:     # second pass with data-table results fed in, so the Options sheet sees bear/base/bull
        dtk = {k: v for k, v in vals.items() if "'!" in k and k.split("!")[0].endswith("ASSUMPTIONS'")
               and re.match(r"[B-F](3[0-2]|3[7-9]|4[01])$", k.split("!")[1])}
        sol2 = xl.calculate(inputs=dtk)
        for k, v in sol2.items():
            if k not in dtk:
                vals[k] = val(v)
        errs = [(k, v) for k, v in vals.items() if isinstance(v, str) and v.startswith("#")]
        OUT["errors"] = errs[:20]
        o = L["opt"]
        OUT["opt"] = {"inputs": {a: vals.get(X("Options", f"B{a}")) for a in range(4, 31)},
                      "cases": {n: [vals.get(X("Options", f"{c}{r}")) for c in "BCDE"] for n, r in
                                (("Bear", 28), ("Base", 29), ("Bull", 30))},
                      "strat": {lab: [vals.get(X("Options", f"{c}{r}")) for c in "BCDEFGHIJ"] for lab, r in o["srows"].items()}}
    OUT["scen"] = {n: {k: g("Assumptions", f"{col}{30 + j}") for k, col in (("ev", "B"), ("eq", "C"), ("ps", "D"), ("up", "E"))}
                   for j, n in enumerate(("Bear", "Base", "Bull"))}
    OUT["sens"] = [[g("Assumptions", f"{cl(2 + i)}{37 + j}") for i in range(5)] for j in range(5)]
    OUT["npv"] = {k: g("NPV", f"B{r}") for k, r in (("pv_fcf", 19), ("tfcf", 20), ("tv", 21), ("pv_tv", 22), ("ev", 23),
                                                    ("eq", 26), ("ps", 28), ("up", 30), ("tv_pct", 31))}
    N_ = len(L["fyears"])
    OUT["fc"] = {k: [g("NPV", f"{cl(3 + i)}{r}") for i in range(N_)] for k, r in
                 (("g", 5), ("m", 6), ("rev", 7), ("oi", 8), ("fcf", 12), ("pv", 16))}
    # price-implied year-10 operating margin (base case, everything else held)
    tgt, M10, PS = L["price"], X("Assumptions", "C13"), X("NPV", "B28")
    f = lambda m: run({M10: m}, PS) - tgt
    lo, hi = -0.5, 1.0
    flo, fhi = f(lo), f(hi)
    if flo > 0:
        OUT["implied_m10"] = "<-50%"
    elif fhi < 0:
        OUT["implied_m10"] = ">100%"
    else:
        for _ in range(30):
            mid = (lo + hi) / 2
            if f(mid) > 0:
                hi = mid
            else:
                lo = mid
        OUT["implied_m10"] = (lo + hi) / 2


def vxml(v):
    if v is None or (isinstance(v, str) and v == ""):
        return None
    if isinstance(v, bool):
        return ' t="b"', "1" if v else "0"
    if isinstance(v, (int, float)):
        return "", repr(float(v)) if isinstance(v, float) else str(v)
    s = str(v)
    if s.startswith("#"):
        return ' t="e"', s
    return ' t="str"', escape(s)


out = {}
for path, sname in sheets.items():
    xml = zin.read(path).decode()
    dtref = set()
    for m in re.finditer(r'<f t="dataTable" ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"', xml):
        for r in range(int(m.group(2)), int(m.group(4)) + 1):
            for c in range(ci(m.group(1)), ci(m.group(3)) + 1):
                dtref.add(f"{cl(c)}{r}")
    seen = set()

    def fix(m):
        attrs, body = m.group(1), m.group(3) or ""
        addr = re.search(r'r="([A-Z]+\d+)"', attrs).group(1)
        is_f, is_dt = "<f" in body, addr in dtref
        if not (is_f or is_dt):
            return m.group(0)
        seen.add(addr)
        vx = vxml(vals.get(X(sname, addr)))
        if vx is None:
            return m.group(0)
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        fpart = re.search(r"<f[^>]*/>|<f[^>]*>.*?</f>", body, re.S)
        return f"<c{attrs}{vx[0]}>{fpart.group(0) if fpart else ''}<v>{vx[1]}</v></c>"

    xml2 = re.sub(r"<c( [^>]*?)(/>|>(.*?)</c>)", lambda m: fix(m) if m.group(2) != "/>" or
                  re.search(r'r="([A-Z]+\d+)"', m.group(1)).group(1) in dtref else m.group(0), xml, flags=re.S)
    miss = dtref - seen
    assert not miss, f"{sname}: data-table cells missing from XML: {sorted(miss)[:5]}"
    out[path] = xml2
# external link rels Id must match <externalBook r:id>
for name in zin.namelist():
    m = re.fullmatch(r"xl/externalLinks/(externalLink\d+)\.xml", name)
    if m:
        rid_ = re.search(r'<externalBook r:id="([^"]+)"', zin.read(name).decode()).group(1)
        rn = f"xl/externalLinks/_rels/{m.group(1)}.xml.rels"
        out[rn] = re.sub(r'Id="[^"]+"', f'Id="{rid_}"', zin.read(rn).decode())
with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zin.infolist():
        data = out[it.filename].encode() if it.filename in out else zin.read(it.filename)
        zo.writestr(it, data, compress_type=zipfile.ZIP_DEFLATED)
json.dump(OUT, open(f"{D}/out/{T}_calc.json", "w"), default=str)
mn = OUT["main"]
print(f"{T}: {L['kind']} {len(vals)} cells {len(errs)} errors {time.time() - t0:.0f}s | value {mn['C11']} "
      f"price {mn['C3']} up {mn['C12']}" + (f" | scen {[round(OUT['scen'][k]['ps'], 2) for k in OUT['scen']]} "
                                            f"impl m10 {OUT['implied_m10']}" if L['kind'] == 'DCF' else ""))
