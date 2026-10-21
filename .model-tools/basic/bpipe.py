"""Pipeline rNPV sheets (Assumptions, Pipeline, NPV) for a pre-revenue / small-revenue company.  Used by bbuild.py.
Programs come from pipeline/programs/<T>.json (ClinicalTrials.gov + curated regulatory overrides); approval odds are
BIO/Informa/QLS 2011-2020 phase-transition rates; peak sales are indication-level placeholders."""
import datetime as dt
from openpyxl.styles import PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.formula import DataTableFormula

HDR = PatternFill("solid", fgColor="D9E1F2")
SEC = PatternFill("solid", fgColor="F2F2F2")
TOP = Border(top=Side(style="thin"))
MM1 = '#,##0.0;(#,##0.0);"-"'
MM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT0 = '0%;(0%);"-"'
USD = '$#,##0.00;($#,##0.00);"-"'
Y0, NY = 2026, 20
YEARS = list(range(Y0, Y0 + NY))
STAGE = {4: "Approved", 3.5: "Filed (NDA/BLA)", 3.0: "Phase 3", 2.5: "Phase 2/3", 2.0: "Phase 2", 1.5: "Phase 1/2", 1.0: "Phase 1"}
RAMP = [0.10, 0.25, 0.45, 0.65, 0.85, 1.00]


def build(wb, T, put, ctx):
    """ctx: name, programs, burn, nol, stub, bsd (date), first_launch.  Returns layout cells."""
    progs = ctx["programs"]
    A = wb.create_sheet("Assumptions")
    PL = wb.create_sheet("Pipeline")
    NP = wb.create_sheet("NPV")
    # ---------------------------------------------------------------- Assumptions
    A.column_dimensions["A"].width = 46
    for col in "BCDEFG":
        A.column_dimensions[col].width = 13
    put(A, "A1", f"{ctx['name']}: pipeline rNPV assumptions", bold=True)
    put(A, "A2", "Blue = input, black = formula, green = link. $mm unless noted.", italic=True)
    put(A, "A4", "Scenario (1 Bear, 2 Base, 3 Bull)", bold=True)
    put(A, "B4", 2, fmt="0")
    put(A, "C4", '=CHOOSE($B$4,"Bear","Base","Bull")', bold=True)
    for r, lab in ((6, "Discount rate override (blank = scenario)"), (7, "Approval-odds multiplier override"),
                   (8, "Peak-sales multiplier override")):
        put(A, f"A{r}", lab)
        put(A, f"B{r}", None, fmt=PCT if r == 6 else "0.00", color="0000FF")
    for i, h in enumerate(["Parameter", "Bear", "Base", "Bull", "Live"]):
        put(A, f"{CL(1 + i)}9", h, bold=True, fill=HDR, align=None if i == 0 else "right")
    params = [(10, "Discount rate", (0.14, 0.12, 0.10), PCT, "$B$6"),
              (11, "Peak-sales multiplier (× indication table)", (0.6, 1.0, 1.4), "0.00", "$B$8"),
              (12, "Approval-odds multiplier (× BIO rates)", (0.8, 1.0, 1.15), "0.00", "$B$7"),
              (13, "Net margin on owned product sales at peak", (0.40, 0.40, 0.40), PCT0, None),
              (14, "Royalty / economics on partnered sales", (0.12, 0.12, 0.12), PCT0, None),
              (15, "Tax rate", (0.21, 0.21, 0.21), PCT0, None)]
    for r, lab, v, fmt, ov in params:
        put(A, f"A{r}", lab)
        for j, x in enumerate(v):
            put(A, f"{CL(2 + j)}{r}", x, fmt=fmt)
        live = f"CHOOSE($B$4,B{r},C{r},D{r})"
        put(A, f"E{r}", f"=IF(ISNUMBER({ov}),{ov},{live})" if ov else f"={live}", fmt=fmt, bold=True)
    put(A, "A17", "Company costs and timing", bold=True, fill=SEC)
    for col in "BCDEFG":
        A[f"{col}17"].fill = SEC
    last_prog_row = 3 + len(progs)
    singles = [(18, "Net operating burn, TTM (operating loss)", round(ctx["burn"], 3), MM1),
               (19, "Years of full burn (until the first launch, max 4)",
                f"=MIN(4,MAX(1,MIN(Pipeline!I4:I{last_prog_row})-{Y0}))", "0"),
               (20, "Burn after that, % of TTM (platform R&D and G&A), if a product is approved", 0.30, PCT0),
               (21, "Tax-loss carryforward at start", round(ctx["nol"], 1), MM),
               (22, "First year, fraction after the balance sheet date", round(ctx["stub"], 4), "0.00"),
               (23, "Balance sheet date", dt.datetime.combine(ctx["bsd"], dt.time()), "yyyy-mm-dd"),
               (25, "Patent / exclusivity life after launch (years)", 12, "0"),
               (26, "Sales erosion per year after exclusivity", 0.50, PCT0)]
    for r, lab, v, fmt in singles:
        put(A, f"A{r}", lab)
        put(A, f"B{r}", v, fmt=fmt)
    put(A, "A16", "Probability at least one owned program is approved")
    put(A, "B16", f"=1-PRODUCT(Pipeline!AG4:AG{last_prog_row})", fmt=PCT0)
    put(A, "A24", "Launch ramp, share of peak in years 1-6")
    for i, v in enumerate(RAMP):
        put(A, f"{CL(2 + i)}24", v, fmt=PCT0)
    # scenario summary + sensitivity (same positions as the DCF sheets, so bcalc reads them the same way)
    put(A, "A27", "Scenario summary", bold=True, fill=SEC)
    heads = [("Pipeline rNPV (ops)", "=NPV!B22", MM), ("Equity value", "=NPV!B25", MM), ("rNPV / share", "=NPV!B27", USD),
             ("Upside", "=NPV!B29", PCT)]
    put(A, "A28", "Scenario", bold=True, fill=HDR)
    put(A, "A29", "Live", italic=True)
    for i, (h, f, fmt) in enumerate(heads):
        put(A, f"{CL(2 + i)}28", h, bold=True, fill=HDR, align="right")
        put(A, f"{CL(2 + i)}29", f, fmt=fmt, italic=True)
        for j in range(3):
            A[f"{CL(2 + i)}{30 + j}"].number_format = fmt
    for j in range(3):
        put(A, f"A{30 + j}", j + 1, fmt='[=1]"Bear";[=2]"Base";"Bull"', color="000000")
    A["B30"] = DataTableFormula(ref="B30:E32", r1="B4", dt2D=False, dtr=False)
    put(A, "A35", "rNPV / share: approval-odds multiplier (rows) x peak-sales multiplier (columns)", bold=True, fill=SEC)
    put(A, "A36", "=NPV!B27", fmt=USD, italic=True)
    pk, ps = [0.5, 0.75, 1.0, 1.5, 2.0], [0.6, 0.8, 1.0, 1.2, 1.4]
    for i, v in enumerate(pk):
        put(A, f"{CL(2 + i)}36", v, fmt="0.00x", bold=True, fill=HDR)
    for j, v in enumerate(ps):
        put(A, f"A{37 + j}", v, fmt="0.00x", bold=True, fill=HDR, align="right")
        for i in range(5):
            A[f"{CL(2 + i)}{37 + j}"].number_format = USD
    A["B37"] = DataTableFormula(ref="B37:F41", r1="B8", r2="B7", dt2D=True)
    # ---------------------------------------------------------------- Pipeline
    PL.column_dimensions["A"].width = 4
    PL.column_dimensions["B"].width = 34
    PL.column_dimensions["C"].width = 28
    for col, w in zip("DEFGHIJKL", (15, 18, 12, 10, 9, 8, 8, 10, 11)):
        PL.column_dimensions[col].width = w
    put(PL, "A1", f"{ctx['name']}: pipeline (risked revenue, $mm)", bold=True)
    put(PL, "A2", "Programs from ClinicalTrials.gov (company as lead sponsor, or partner of an industry sponsor), one row "
                  "per indication; approval odds from BIO 2011-2020 phase-transition rates.", italic=True)
    hd = ["No.", "Program (indication)", "Drug(s)", "Stage", "Area / rate class", "Economics", "Peak", "PoS", "Launch",
          "LOE", "Share", "Risked PV"]
    for i, h in enumerate(hd):
        put(PL, f"{CL(1 + i)}3", h, bold=True, fill=HDR, align="center" if i else None)
    YC = {y: CL(13 + i) for i, y in enumerate(YEARS)}
    put(PL, "AG3", "1 - PoS (owned)", bold=True, fill=HDR)
    PL.column_dimensions["AG"].width = 13
    for y in YEARS:
        put(PL, f"{YC[y]}3", y, fmt="0", bold=True, fill=HDR)
        PL.column_dimensions[YC[y]].width = 9
    ly = YC[YEARS[-1]]
    for k, p in enumerate(progs):
        r = 4 + k
        put(PL, f"A{r}", k + 1, fmt="0", color="000000")
        put(PL, f"B{r}", p["bucket"])
        put(PL, f"C{r}", ", ".join(p.get("drugs", [])[:2])[:60])
        put(PL, f"D{r}", STAGE.get(p["stage"], str(p["stage"])))
        put(PL, f"E{r}", f"{p['area']} / {p.get('rate_class', 'area')}")
        put(PL, f"F{r}", "Owned" if p["owned"] else "Partnered")
        put(PL, f"G{r}", p["peak"], fmt=MM)
        put(PL, f"H{r}", p["pos"], fmt=PCT)
        put(PL, f"I{r}", p["launch"], fmt="0")
        put(PL, f"J{r}", f"=I{r}+Assumptions!$B$25", fmt="0")
        share = p.get("share_override")
        if share is None and p.get("share") not in (0.40, 0.12):
            share = p["share"]
        put(PL, f"K{r}", share if share is not None else ("=Assumptions!$E$13" if p["owned"] else "=Assumptions!$E$14"),
            fmt=PCT0)
        for y in YEARS:
            c = YC[y]
            f = (f"=IF({c}$3<$I{r},0,$G{r}*Assumptions!$E$11*INDEX(Assumptions!$B$24:$G$24,MIN(6,{c}$3-$I{r}+1))"
                 f"*IF({c}$3>$J{r},(1-Assumptions!$B$26)^({c}$3-$J{r}),1)*IF($H{r}>=1,1,MIN(0.97,$H{r}*Assumptions!$E$12)))")
            put(PL, f"{c}{r}", f, fmt=MM1)
        put(PL, f"L{r}", f"=K{r}*SUMPRODUCT(M{r}:{ly}{r},NPV!$C$15:$V$15)", fmt=MM1)
        put(PL, f"AG{r}", f'=IF(F{r}="Owned",1-IF(H{r}>=1,1,MIN(0.97,H{r}*Assumptions!$E$12)),1)', fmt=PCT0)
    tr = 4 + len(progs)
    put(PL, f"B{tr}", "Risked revenue", bold=True)
    put(PL, f"B{tr + 1}", "Risked contribution (revenue x share)", bold=True)
    put(PL, f"L{tr}", f"=SUM(L4:L{tr - 1})", fmt=MM1, bold=True, border=TOP)
    for y in YEARS:
        c = YC[y]
        put(PL, f"{c}{tr}", f"=SUM({c}4:{c}{tr - 1})", fmt=MM1, bold=True, border=TOP)
        put(PL, f"{c}{tr + 1}", f"=SUMPRODUCT({c}4:{c}{tr - 1},$K$4:$K${tr - 1})", fmt=MM1, bold=True)
    # ---------------------------------------------------------------- NPV
    NP.column_dimensions["A"].width = 40
    NP.column_dimensions["B"].width = 13
    put(NP, "A1", f"{ctx['name']}: risk-adjusted NPV ($mm)", bold=True)
    NC = {y: CL(3 + i) for i, y in enumerate(YEARS)}
    labels = {3: "Year", 4: "Risked pipeline revenue", 5: "Risked contribution", 6: "Corporate burn", 7: "Pre-tax cash flow",
              8: "Tax-loss carryforward used", 9: "Taxable income", 10: "Tax", 11: "Free cash flow", 12: "Year fraction",
              13: "Discount time (mid-year)", 14: "Discount factor", 15: "PV factor", 16: "PV of free cash flow",
              17: "Tax-loss carryforward, end of year"}
    for r, lab in labels.items():
        put(NP, f"A{r}", lab, bold=r in (3, 11, 16))
    for k, y in enumerate(YEARS):
        c, p = NC[y], (NC[y - 1] if k else None)
        NP.column_dimensions[c].width = 9
        put(NP, f"{c}3", y, fmt="0", bold=True, fill=HDR, color="000000")
        put(NP, f"{c}4", f"=Pipeline!{YC[y]}{tr}", fmt=MM1)
        put(NP, f"{c}5", f"=Pipeline!{YC[y]}{tr + 1}", fmt=MM1)
        put(NP, f"{c}6", f"=-IF({c}3-{Y0}<Assumptions!$B$19,Assumptions!$B$18,Assumptions!$B$18*Assumptions!$B$20*Assumptions!$B$16)",
            fmt=MM1)
        put(NP, f"{c}7", f"={c}5+{c}6", fmt=MM1)
        prev_pool = "Assumptions!$B$21" if not p else f"{p}17"
        put(NP, f"{c}8", f"=MIN({prev_pool},MAX(0,{c}7))", fmt=MM1)
        put(NP, f"{c}9", f"=MAX(0,{c}7)-{c}8", fmt=MM1)
        put(NP, f"{c}10", f"={c}9*Assumptions!$E$15", fmt=MM1)
        put(NP, f"{c}11", f"={c}7-{c}10", fmt=MM1, bold=True)
        put(NP, f"{c}12", "=Assumptions!$B$22" if not p else 1, fmt="0.00", color=None if not p else "000000")
        put(NP, f"{c}13", f"={c}12/2" if not p else f"={p}13+{p}12/2+{c}12/2", fmt="0.00")
        put(NP, f"{c}14", f"=1/(1+Assumptions!$E$10)^{c}13", fmt="0.000")
        put(NP, f"{c}15", f"={c}12*{c}14", fmt="0.000")
        put(NP, f"{c}16", f"={c}11*{c}15", fmt=MM1, bold=True)
        put(NP, f"{c}17", f"={prev_pool}-{c}8+MAX(0,-{c}7)", fmt=MM1)
    LC = NC[YEARS[-1]]
    summ = [(19, "PV of risked pipeline contribution", f"=SUMPRODUCT(C5:{LC}5,C15:{LC}15)", MM1),
            (20, "PV of corporate burn", f"=SUMPRODUCT(C6:{LC}6,C15:{LC}15)", MM1),
            (21, "PV of taxes", f"=-SUMPRODUCT(C10:{LC}10,C15:{LC}15)", MM1),
            (22, "Pipeline rNPV (operations)", f"=SUM(C16:{LC}16)", MM1),
            (23, "Cash & investments", "=Main!C7", MM1),
            (24, "Debt", "=Main!C8", MM1),
            (25, "Equity value", "=B22+B23-B24", MM1),
            (26, "Shares (mm)", "=Main!C4", MM1),
            (27, "rNPV / share", "=B25/B26", USD),
            (28, "Price", "=Main!C3", USD),
            (29, "Upside", "=B27/B28-1", PCT)]
    for r, lab, f, fmt in summ:
        put(NP, f"A{r}", lab, bold=r in (22, 25, 27))
        put(NP, f"B{r}", f, fmt=fmt, bold=r in (22, 25, 27), border=TOP if r in (22, 25, 27) else None)
    return {"pk": pk, "ps": ps, "n_programs": len(progs), "prog_rows": [4, tr - 1], "years": YEARS}
