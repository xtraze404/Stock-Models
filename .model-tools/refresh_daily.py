#!/usr/bin/env python3
"""Daily post-close refresh of the Stock Models, ending with a paper-trader run.

usage:
  venv/bin/python refresh_daily.py [--date YYYY-MM-DD] [--dry-run] [--accept-changes] [--from STEP] [--to STEP]
                                   [--refetch] [--no-trade] [--force-trade] [--jobs 6]

Steps (names for --from / --to):
  1 closes    Nasdaq regular-session closes -> closes_<yyyymmdd>.json; missing names re-fetched at 3 threads,
              then from the historical feed
  2 splits    moves over 60% checked against Nasdaq's "Stock Split : <date>" notice and its split-adjusted previous
              close; basic/templates.json C4 adjusted (like REVB 1-for-7 and MBRX 1-for-20 on 2026-10-20)
  3 prices    prices into basic/tickers.json and basic/screener_cache.json; bbuild + bcalc for basic/todo.txt;
              reprice.py for the 16 hand-built models
  4 validate  every model's NPV/share (Main C11) must equal the installed workbook's, except split-adjusted names;
              stops before installing otherwise (override with --accept-changes)
  5 notes     bnote / bhub / brep with the close date
  6 install   workbooks -> Pharma/, notes -> Vault; skips files with a ~$ lock or edited since the last install
  7 screener  Pharma F prices + F1 "Close YYYY-MM-DD", bscreener (layout v2: J/K/L), fixlinks; installs only if
              nothing but column F changed. Skipped while ~$Screener.xlsx exists (Excel has it open)
  8 recs      universe.json -> rec.py -> rec_sheet.py + fixlinks; installs only if nothing but the Recommendations
              sheet changed
  9 trade     trader/paper_trader.py run --prices closes_<yyyymmdd>.json (only when the Screener is at this close)

Paths come from this file's location (.../Stock Models/.model-tools), so a full copy of the folder runs against the
copy. The vault is ../Vault next to "Stock Models" unless SM_VAULT is set.
Each run logs to runs/<yyyymmdd>/ (refresh.log, build logs, staged Screener files, state.json for --from).
Exit codes: 0 done, 2 stopped by a safety check (nothing after that step ran), 1 a script failed."""
import argparse, ast, concurrent.futures as cf, datetime as dt, glob, hashlib, json, os, re, shutil, subprocess, sys
import time, urllib.request
from zoneinfo import ZoneInfo

TOOLS = os.path.dirname(os.path.abspath(__file__))
BASIC = f"{TOOLS}/basic"
ROOT = os.path.dirname(TOOLS)                                  # .../Stock Models
PHARMA = f"{ROOT}/Pharma"
SCREENER = f"{ROOT}/Screener.xlsx"
VAULT = os.environ.get("SM_VAULT") or os.path.normpath(f"{ROOT}/../Vault")
COMPANIES = f"{VAULT}/Finance/Companies"
GALLERY = f"{VAULT}/Gallery/Image"
HUB = (f"{BASIC}/notes_hub/Basic NPV Models.md", f"{VAULT}/Finance/Basic NPV Models.md")
REPORT = (f"{BASIC}/report/Large-Cap Upside Review.md", f"{VAULT}/Finance/Large-Cap Upside Review.md")
TRADER = f"{TOOLS}/trader"
PY = f"{TOOLS}/venv/bin/python"
MANIFEST = f"{TOOLS}/install_manifest.json"     # md5 of every file at its last install ("edited since" check)
SPLITS = f"{BASIC}/splits_applied.json"         # splits already applied to templates.json (and the trader)
TICKERS, CACHE, TEMPLATES = f"{BASIC}/tickers.json", f"{BASIC}/screener_cache.json", f"{BASIC}/templates.json"
UNIVERSE = f"{TOOLS}/universe.json"
# hand-built models: repriced in place, never rebuilt (kind = universe.json / rec.py model type)
BUILT = {"LLY": "Full DCF", "EVMN": "rNPV", **{t: "PDUFA rNPV" for t in (
    "INO", "AGIO", "CYTK", "SMMT", "SVRA", "CAPR", "BBIO", "PYPD", "COGT", "ANAB", "VNDA", "CORT", "MLYS", "PRAX")}}
ALIASES = {"QNCX": "IRLA"}                      # QNCX became IRulya Therapeutics (IRLA) on 2026-10-09
SPLIT_MOVE = 0.60                               # a move this large is checked for a split
MAX_MISSING = 0.05                              # stop if more than this share of tickers has no close
NY = ZoneInfo("America/New_York")
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 "
                    "Safari/537.36", "Accept": "application/json"}
STEPS = ["closes", "splits", "prices", "validate", "notes", "install", "screener", "recs", "trade"]
TITLES = {"closes": "Closes", "splits": "Split check", "prices": "Prices and rebuild", "validate": "Validation",
          "notes": "Notes", "install": "Install", "screener": "Screener", "recs": "Recommendations",
          "trade": "Paper trader"}
LOG = None


class Stop(Exception):
    """A safety check failed (code 2) or a script failed (code 1). Steps after this one don't run."""

    def __init__(self, msg, code=2):
        super().__init__(msg)
        self.code = code


# ---------------------------------------------------------------- helpers
def say(*a):
    s = " ".join(str(x) for x in a)
    print(s, flush=True)
    if LOG:
        LOG.write(s + "\n")
        LOG.flush()


def rel(p):
    return os.path.relpath(p, os.path.dirname(ROOT))


def md5(p):
    if not os.path.exists(p):
        return None
    h = hashlib.md5()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def jload(p, default=None):
    if default is not None and not os.path.exists(p):
        return default
    with open(p) as f:
        return json.load(f)


def jsave(p, obj, **kw):
    tmp = p + ".tmp"
    with open(tmp, "w") as f:
        json.dump(obj, f, **kw)
    os.replace(tmp, p)


def lock_of(path):
    return os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))


def is_locked(path):
    return os.path.exists(lock_of(path))


def copy_atomic(src, dst):
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    tmp = os.path.join(os.path.dirname(dst), "." + os.path.basename(dst) + ".refresh-tmp")
    shutil.copyfile(src, tmp)
    os.replace(tmp, dst)


def run(cmd, cwd, log=None, echo=False, env=None):
    """Run a pipeline script; stop the refresh if it fails. Returns stdout."""
    r = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, env=env)
    out = r.stdout + (("\n" + r.stderr) if r.stderr.strip() else "")
    if log:
        with open(log, "a") as f:
            f.write(f"$ {' '.join(cmd)}\n{out}\n")
    if echo and out.strip():
        for ln in out.rstrip().splitlines():
            say("   ", ln)
    if r.returncode:
        tail = "\n".join(out.strip().splitlines()[-15:])
        raise Stop(f"{os.path.basename(cmd[1])} failed (exit {r.returncode}):\n{tail}", 1)
    return r.stdout


def short(cmd):
    """A command for printing: script name + args, ticker lists shortened."""
    parts = [os.path.basename(cmd[1])] + [rel(c) if c.startswith("/") else c for c in cmd[2:]]
    if len(parts) > 12:
        parts = parts[:8] + [f"... ({len(parts) - 8} more)"]
    return " ".join(f'"{p}"' if " " in p else p for p in parts)


def main_vals(path, cells=("C3", "C4", "C11", "C12")):
    """Cached values of Main!<cells> read straight from the xlsx XML (fast; no openpyxl)."""
    import zipfile
    out = {a: None for a in cells}
    if not os.path.exists(path):
        return out
    z = zipfile.ZipFile(path)
    wbx, rels = z.read("xl/workbook.xml").decode(), z.read("xl/_rels/workbook.xml.rels").decode()
    m = re.search(r'<sheet [^>]*name="Main"[^>]*r:id="(rId\d+)"', wbx) or \
        re.search(r'<sheet [^>]*r:id="(rId\d+)"[^>]*name="Main"', wbx)
    rid = m.group(1)
    m = re.search(rf'<Relationship [^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
    t = m.group(1).lstrip("/")
    xml = z.read(t if t.startswith("xl/") else "xl/" + t).decode()
    for a in cells:
        m = re.search(rf'<c r="{a}"([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S)
        v = re.search(r"<v>(.*?)</v>", m.group(2), re.S) if m and m.group(2) else None
        if v:
            ty = re.search(r'\bt="(\w+)"', m.group(1))
            out[a] = v.group(1) if ty and ty.group(1) in ("str", "e", "s", "inlineStr") else float(v.group(1))
    return out


def same(a, b):
    if a in ("", None) and b in ("", None):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return a == b


def todo_list():
    return [ln.strip() for ln in open(f"{BASIC}/todo.txt") if ln.strip()]


def user_notes():
    """bnote.py's USER_NOTES: notes the user wrote; bnote merges into the vault copy, so edits carry forward."""
    m = re.search(r"^USER_NOTES = (\{[^}]*\})", open(f"{BASIC}/bnote.py").read(), re.M)
    return ast.literal_eval(m.group(1)) if m else set()


def report_image():
    """Image number the Large-Cap report embeds (kept fixed so the vault link never changes)."""
    for p in (REPORT[1], REPORT[0]):
        if os.path.exists(p):
            m = re.search(r"Image-(\d{5})\.png", open(p).read())
            if m:
                return int(m.group(1))
    pngs = sorted(glob.glob(f"{BASIC}/report/Image-*.png"))
    if pngs:
        return int(re.search(r"Image-(\d+)", pngs[-1]).group(1))
    raise Stop("Can't find the Large-Cap report's image number")


def image_dst(num):
    name = f"Image-{num:05d}.png"
    hit = glob.glob(f"{GALLERY}/*/{name}")
    if hit:
        return hit[0]
    lo = (num - 1) // 100 * 100 + 1
    return f"{GALLERY}/IMG {lo:05d}-{lo + 99:05d}/{name}"


def install_pairs():
    """(source, installed path, kind). 'regen' files are rebuilt from scratch, so a file edited since the last install
    is kept and skipped. 'carry' files were read from the installed copy this run (hand-built models, user notes), so
    they're skipped only if the installed copy changed after it was read."""
    un, todo, img = user_notes(), todo_list(), report_image()
    P = [(f"{BASIC}/cached/{t}.xlsx", f"{PHARMA}/{t}.xlsx", "regen") for t in todo]
    P += [(f"{TOOLS}/reprice_out/{t}.xlsx", f"{PHARMA}/{t}.xlsx", "carry") for t in BUILT]
    P += [(f"{BASIC}/notes/{t}.md", f"{COMPANIES}/{t}.md", "carry" if t in un else "regen") for t in todo]
    P += [HUB + ("regen",), REPORT + ("regen",), (f"{BASIC}/report/Image-{img:05d}.png", image_dst(img), "regen")]
    return P


def key(p):
    return rel(p)


def load_manifest(ctx):
    if os.path.exists(MANIFEST):
        return jload(MANIFEST)
    if ctx.get("manifest") is not None:
        return ctx["manifest"]
    raise Stop(f"No install record ({rel(MANIFEST)}). Run from step 1-3 once so it can be created.")


def ensure_manifest(ctx):
    """First run: record the last pipeline outputs as the installed state (before step 3 overwrites them)."""
    if os.path.exists(MANIFEST):
        return
    M = {key(dst): md5(src) for src, dst, _ in install_pairs() if os.path.exists(src)}
    same_ = sum(1 for src, dst, _ in install_pairs() if os.path.exists(dst) and md5(dst) == M.get(key(dst)))
    say(f"No install record yet: {'would create' if ctx['dry'] else 'created'} one from the last outputs "
        f"({same_} of {len(M)} installed files match them; any that don't count as edited by you).")
    ctx["manifest"] = M
    if not ctx["dry"]:
        jsave(MANIFEST, M, indent=0, sort_keys=True)


def edited(dst, M):
    cur = md5(dst)
    return cur is not None and cur != M.get(key(dst))


def load_closes(ctx):
    p = ctx["closes_file"]
    if not os.path.exists(p):
        raise Stop(f"{rel(p)} doesn't exist yet; run the closes step first.")
    return {t: v["px"] for t, v in jload(p).items() if isinstance(v, dict) and v.get("px")}


def nasdaq(path):
    for attempt in range(3):
        try:
            return json.load(urllib.request.urlopen(urllib.request.Request(
                f"https://api.nasdaq.com/api/quote/{path}", headers=UA), timeout=20)).get("data") or {}
        except Exception:
            time.sleep(1.5 * (attempt + 1))
    return {}


def parse_day(s):
    s = s.strip()
    for f in ("%b %d, %Y", "%m/%d/%Y", "%Y-%m-%d", "%B %d, %Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def simple_ratio(f):
    """Price factor (split-adjusted previous close / our previous close) as num/den with den <= 10."""
    for den in range(1, 11):
        num = round(f * den)
        if num >= 1 and (num, den) != (1, 1) and abs(num / den / f - 1) < 0.03:
            return num, den
    return None


def split_label(num, den):
    if den == 1:
        return f"1-for-{num} reverse"
    if num == 1:
        return f"{den}-for-1"
    return f"{den}-for-{num}"


# ---------------------------------------------------------------- 1 closes
def step_closes(ctx):
    T = jload(TICKERS)
    tick = [x["t"] for x in T]
    out = ctx["closes_file"]
    prev_close = jload(CACHE).get("close")
    if prev_close and ctx["iso"] < prev_close:
        raise Stop(f"--date {ctx['iso']} is before the close already loaded ({prev_close}); that would roll prices back.")
    have = {}
    if os.path.exists(out) and not ctx["args"].refetch:
        have = {t: v for t, v in jload(out).items() if t in tick and isinstance(v, dict) and v.get("px")}
    need = [t for t in tick if t not in have]
    mode = "quote feed ('Closed at ...'), then historical feed" if ctx["live"] else "historical feed (official close)"
    say(f"{len(tick)} tickers; {len(have)} already in {rel(out)}; {len(need)} to fetch from the {mode}")
    if ctx["dry"]:
        if need:
            cmd = [PY, f"{TOOLS}/closes_live.py", ctx["stamp"], ctx["iso"], out] + need if ctx["live"] else \
                [PY, f"{TOOLS}/closes.py", ctx["date"].strftime("%m/%d/%Y"), out] + need
            say("  would run:", short(cmd))
        return

    def merge(p, src_default=None):
        for t, v in jload(p).items():
            px = v.get("px") if isinstance(v, dict) else v
            if px:
                have[t] = {"px": px, "src": (v.get("src") if isinstance(v, dict) else None) or src_default}

    log = f"{ctx['run_dir']}/closes.log"
    if need and ctx["live"]:
        for i in (1, 2):                          # second pass: re-fetch the misses (3 threads, like the first)
            tmp = f"{ctx['run_dir']}/closes_live_{i}.json"
            say(f"  {'fetching' if i == 1 else 're-fetching'} {len(need)} from the quote feed")
            run([PY, "closes_live.py", ctx["stamp"], ctx["iso"], tmp] + need, cwd=TOOLS, log=log, echo=True)
            merge(tmp)
            need = [t for t in tick if t not in have]
            if not need:
                break
    if need:
        tmp = f"{ctx['run_dir']}/closes_hist.json"
        say(f"  fetching {len(need)} from the historical feed")
        run([PY, "closes.py", ctx["date"].strftime("%m/%d/%Y"), tmp] + need, cwd=TOOLS, log=log, echo=True)
        merge(tmp, "nasdaq-hist")
        need = [t for t in tick if t not in have]
    fetched = any(os.path.exists(f"{ctx['run_dir']}/{f}") for f in ("closes_live_1.json", "closes_hist.json"))
    if fetched or not os.path.exists(out):
        jsave(out, {t: have.get(t, {"px": None, "src": None}) for t in tick}, indent=0)
    ctx["state"]["missing_closes"] = need
    say(f"  {len(have)} closes in {rel(out)}" + (f"; no close for {need}" if need else ""))
    if len(need) > MAX_MISSING * len(tick):
        raise Stop(f"{len(need)} of {len(tick)} tickers have no close for {ctx['iso']}. Market holiday, or the feed "
                   "isn't posting closes yet. Nothing has been changed.")
    if need:
        say(f"  those keep their previous price")
    prev = {x["t"]: x["price"] for x in T}
    if prev_close and prev_close != ctx["iso"]:
        unchanged = sum(1 for t, v in have.items() if prev.get(t) == v["px"])
        if unchanged > 0.5 * len(have):
            raise Stop(f"{unchanged} of {len(have)} closes equal the {prev_close} prices; the feed looks stale. "
                       f"Delete {rel(out)} or use --refetch once it updates.")


# ---------------------------------------------------------------- 2 splits
def step_splits(ctx):
    if not os.path.exists(ctx["closes_file"]):
        say("no closes file yet (dry run): the split check runs on the fetched closes")
        return
    px = load_closes(ctx)
    prev = {x["t"]: x["price"] for x in jload(TICKERS)}
    prev_close = jload(CACHE).get("close")
    applied = jload(SPLITS, {})
    tpl = jload(TEMPLATES)
    movers = [(t, prev[t], px[t]) for t in prev if px.get(t) and prev.get(t) and abs(px[t] / prev[t] - 1) > SPLIT_MOVE]
    say(f"{len(movers)} moves over {SPLIT_MOVE:.0%} vs. the {prev_close} prices" +
        (": " + ", ".join(f"{t} {b / a - 1:+.0%}" for t, a, b in movers) if movers else ""))
    splits, manual, changed = ctx["state"].setdefault("splits", {}), [], False
    for t, old, new in movers:
        sym = ALIASES.get(t, t)
        info = nasdaq(f"{sym}/info?assetclass=stocks")
        notices = [parse_day(m.group(1)) for n in info.get("notifications") or [] for e in n.get("eventTypes") or []
                   for m in [re.search(r"Stock Split\s*:\s*(.+)", e.get("message") or "")] if m]
        sd = next((d for d in notices if d and (not prev_close or prev_close < d.isoformat() <= ctx["iso"])), None)
        if not sd:
            say(f"  {t}: {new / old - 1:+.0%} with no split notice dated after {prev_close}"
                + (f" (notice for {', '.join(map(str, notices))})" if notices else "") + "; price used as is")
            continue
        summ = (nasdaq(f"{sym}/summary?assetclass=stocks").get("summaryData") or {}).get("PreviousClose") or {}
        try:
            adj = float(str(summ.get("value")).replace("$", "").replace(",", ""))
        except ValueError:
            adj = None
        rn = simple_ratio(adj / old) if adj else None
        if not rn or not 0.4 < (new / old) / (rn[0] / rn[1]) < 2.5:
            manual.append(f"{t}: split notice {sd} but the ratio isn't clear (previous close {old}, Nasdaq's adjusted "
                          f"previous close {adj}, close {new})")
            continue
        num, den = rn
        label = split_label(num, den)
        k = f"{t}:{sd.isoformat()}"
        if k in applied:
            splits[t] = applied[k]
            say(f"  {t}: {label} split effective {sd} already applied ({applied[k]['new_c4']})")
            continue
        if t in BUILT:
            manual.append(f"{t}: {label} split effective {sd}. Hand-built model: divide Main C4 by {num / den:g} in "
                          f"Pharma/{t}.xlsx yourself, then run again")
            continue
        c4 = tpl[t]["C4"]
        if c4 is None:
            manual.append(f"{t}: {label} split effective {sd}, but templates.json has no C4 (ADS count from bbuild)")
            continue
        expr = c4[1:] if isinstance(c4, str) and c4.startswith("=") else repr(c4)
        c4n = f"=({expr})/{num}" if den == 1 else f"=({expr})*{den}" if num == 1 else f"=({expr})*{den}/{num}"
        rec = dict(date=sd.isoformat(), num=num, den=den, label=label, old_c4=c4, new_c4=c4n, prev=old, adj_prev=adj,
                   close=new, run=ctx["iso"], trader=False)
        say(f"  {t}: {label} split effective {sd} (previous close {old} -> {adj} split-adjusted). "
            f"templates.json C4 {c4!r} -> {c4n!r}" + (" [dry run: not written]" if ctx["dry"] else ""))
        if ctx["dry"]:
            continue
        tpl[t]["C4"] = c4n
        tpl[t]["C4c"] = (tpl[t].get("C4c") or "").rstrip() + \
            f" Adjusted for a {label} split effective {sd.isoformat()} (Nasdaq split notice)."
        applied[k] = splits[t] = rec
        changed = True
    if changed:
        shutil.copyfile(TEMPLATES, f"{ctx['run_dir']}/templates.before-splits.json")
        jsave(TEMPLATES, tpl, indent=0)
        jsave(SPLITS, applied, indent=1, sort_keys=True)
    if manual:
        raise Stop("Split check needs you:\n  " + "\n  ".join(manual))


# ---------------------------------------------------------------- 3 prices
def step_prices(ctx):
    have = os.path.exists(ctx["closes_file"])
    px = load_closes(ctx) if have else {}
    T, SC, todo = jload(TICKERS), jload(CACHE), todo_list()
    n = 0
    for x in T:
        if x["t"] in px and x["price"] != px[x["t"]]:
            x["price"] = px[x["t"]]
            n += 1
    for row in SC["rows"]:
        if row[1] in px and row[1] != "Ticker":
            row[2] = px[row[1]]
    SC["close"] = ctx["iso"]
    say(f"{n} prices changed in basic/tickers.json and basic/screener_cache.json (close {ctx['iso']})")
    build = [[PY, "bbuild.py", "T"], [PY, "bcalc.py", "T"]]
    if ctx["dry"]:
        say(f"  would run for {len(todo)} tickers in basic/todo.txt ({ctx['args'].jobs} at a time): "
            + " && ".join(short(c) for c in build))
        for t in BUILT:
            say(f"  would run: reprice.py Pharma/{t}.xlsx {px.get(t, 'PRICE')} reprice_out/{t}.xlsx")
        return
    jsave(TICKERS, T, indent=0)
    jsave(CACHE, SC)
    logs = f"{ctx['run_dir']}/build"
    os.makedirs(logs, exist_ok=True)

    def one(t):
        with open(f"{logs}/{t}.txt", "w") as fh:
            for c in build:
                if subprocess.run([x if x != "T" else t for x in c], cwd=BASIC, stdout=fh,
                                  stderr=subprocess.STDOUT).returncode:
                    return t, c[1]
        return t, None

    t0, fails = time.time(), []
    with cf.ThreadPoolExecutor(ctx["args"].jobs) as ex:
        for i, (t, err) in enumerate(ex.map(one, todo), 1):
            if err:
                fails.append(f"{t} ({err})")
            if i % 100 == 0 or i == len(todo):
                say(f"  built {i}/{len(todo)} ({time.time() - t0:.0f}s)")
    if fails:
        raise Stop(f"Build failed for {len(fails)}: {fails[:10]}. Logs: {rel(logs)}/", 1)
    # hand-built models: reprice the installed workbook (user edits carry forward)
    prices = {x["t"]: x["price"] for x in T}
    carry = ctx["state"].setdefault("carry", {})
    for t in BUILT:
        carry[key(f"{PHARMA}/{t}.xlsx")] = md5(f"{PHARMA}/{t}.xlsx")
        if is_locked(f"{PHARMA}/{t}.xlsx"):
            say(f"  {t}.xlsx is open in Excel; repricing the saved copy, install will skip it")
    os.makedirs(f"{TOOLS}/reprice_out", exist_ok=True)

    def rp(t):
        o = run([PY, "reprice.py", f"{PHARMA}/{t}.xlsx", repr(float(prices[t])), f"{TOOLS}/reprice_out/{t}.xlsx"],
                cwd=TOOLS, log=f"{logs}/reprice_{t}.txt")
        return t, json.loads(o.strip().splitlines()[-1])

    with cf.ThreadPoolExecutor(ctx["args"].jobs) as ex:
        res = dict(ex.map(rp, BUILT))
    say(f"  repriced {len(res)} hand-built models: " +
        ", ".join(f"{t} {r['C11']:.2f}" + (f" ({r['errors']} errors)" if r["errors"] else "") for t, r in res.items()))


# ---------------------------------------------------------------- 4 validate
def step_validate(ctx):
    M = load_manifest(ctx)
    splits = set(ctx["state"].get("splits", {}))
    rows, skipped_edit, errs = [], [], []
    for t in todo_list():
        dst = f"{PHARMA}/{t}.xlsx"
        if edited(dst, M):
            skipped_edit.append(t)
            continue
        C = jload(f"{BASIC}/out/{t}_calc.json")
        if C.get("errors"):
            errs.append(f"{t} ({len(C['errors'])})")
        rows.append((t, main_vals(dst, ("C11",))["C11"], C["main"]["C11"]))
    for t in BUILT:
        rows.append((t, main_vals(f"{PHARMA}/{t}.xlsx", ("C11",))["C11"],
                     main_vals(f"{TOOLS}/reprice_out/{t}.xlsx", ("C11",))["C11"]))
    diff = [(t, a, b) for t, a, b in rows if not same(a, b)]
    exempt = [d for d in diff if d[0] in splits]
    bad = [d for d in diff if d[0] not in splits]
    say(f"{len(rows) - len(diff)} of {len(rows)} models have the installed NPV/share (Main C11)"
        + (" [dry run: current outputs vs. installed]" if ctx["dry"] else ""))
    fmt = lambda d: f"{d[0]} {d[1]!r} -> {d[2]!r}" + (f" ({d[2] / d[1] - 1:+.1%})" if all(
        isinstance(x, float) for x in d[1:]) and d[1] else "")
    for d in exempt:
        say(f"  split-adjusted (allowed): {fmt(d)}")
    if skipped_edit:
        say(f"  not compared, edited since the last install (install will keep your version): {skipped_edit}")
    if errs:
        say(f"  formula errors in {len(errs)} rebuilt models: {errs[:10]}")
    if bad:
        for d in bad[:40]:
            say(f"  CHANGED: {fmt(d)}")
        if len(bad) > 40:
            say(f"  ... and {len(bad) - 40} more")
        if ctx["args"].accept_changes:
            say(f"  {len(bad)} changed; continuing because of --accept-changes")
        elif ctx["dry"]:
            say(f"  a real run would stop here unless --accept-changes")
        else:
            raise Stop(f"{len(bad)} models' NPV/share changed on a price update, so something other than the price "
                       f"changed. Nothing was installed. If the changes are expected, run again with "
                       f"--from validate --accept-changes.")


# ---------------------------------------------------------------- 5 notes
def step_notes(ctx):
    todo, img = todo_list(), report_image()
    jobs = ctx["args"].jobs
    chunks = [todo[i::jobs] for i in range(jobs)]
    cmds = [[PY, "bnote.py", "--close", ctx["iso"]] + c for c in chunks if c]
    say(f"{len(todo)} notes (bnote.py, {len(cmds)} batches), hub (bhub.py), Large-Cap report (brep.py, Image-{img:05d})"
        f", close {ctx['iso']}")
    if ctx["dry"]:
        say("  would run:", short(cmds[0]), "| ...")
        say("  would run:", f"bhub.py --close {ctx['iso']}", "|", f"brep.py {img} --close {ctx['iso']}")
        return
    carry = ctx["state"].setdefault("carry", {})
    for t in user_notes():
        carry[key(f"{COMPANIES}/{t}.md")] = md5(f"{COMPANIES}/{t}.md")
    log = f"{ctx['run_dir']}/notes.log"
    env = dict(os.environ, SM_VAULT=VAULT)
    with cf.ThreadPoolExecutor(jobs) as ex:
        list(ex.map(lambda c: run(c, cwd=BASIC, log=log, env=env), cmds))
    run([PY, "bhub.py", "--close", ctx["iso"]], cwd=BASIC, log=log, echo=True)
    run([PY, "brep.py", str(img), "--close", ctx["iso"]], cwd=BASIC, log=log, echo=True)


# ---------------------------------------------------------------- 6 install
def step_install(ctx):
    M = dict(load_manifest(ctx))
    carry = ctx["state"].get("carry", {})
    done, same_, locked, edits, moved, missing = [], 0, [], [], [], []
    for src, dst, kind in install_pairs():
        k = key(dst)
        if not os.path.exists(src):
            missing.append(rel(src))
            continue
        if dst.endswith(".xlsx") and is_locked(dst):
            locked.append(os.path.basename(dst))
            continue
        cur, new = md5(dst), md5(src)
        if kind == "regen" and cur is not None and cur != M.get(k):
            edits.append(rel(dst))
            continue
        if kind == "carry" and cur != carry.get(k, M.get(k, cur)):
            moved.append(rel(dst))
            continue
        if cur == new:
            same_ += 1
            M[k] = new
            continue
        done.append(rel(dst))
        if not ctx["dry"]:
            copy_atomic(src, dst)
            M[k] = new
    verb = "would install" if ctx["dry"] else "installed"
    say(f"{verb} {len(done)} files, {same_} already up to date")
    if done:
        xl = sum(1 for p in done if p.endswith(".xlsx"))
        say(f"  {xl} workbooks, {len(done) - xl} notes/report files" + (f", e.g. {done[:3]}" if done else ""))
    for lab, xs in (("skipped, open in Excel (~$ lock)", locked), ("skipped, edited since the last install", edits),
                    ("skipped, changed while the refresh ran", moved), ("no output to install", missing)):
        if xs:
            say(f"  {lab}: {len(xs)}: {xs[:12]}" + (" ..." if len(xs) > 12 else ""))
    ctx["state"]["install_skipped"] = locked + edits + moved
    if not ctx["dry"]:
        jsave(MANIFEST, M, indent=0, sort_keys=True)


# ---------------------------------------------------------------- 7 screener
def wb_snapshot(path, skip=()):
    """Everything in a workbook this pipeline could disturb: sheet order, links, names, and per sheet the merged
    ranges, conditional-format ranges, filter, freeze panes and each cell's value, number format and basic style."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    meta = {"sheets": wb.sheetnames, "links": [l.file_link.Target for l in wb._external_links],
            "names": sorted(wb.defined_names.keys())}
    cells = {}
    for ws in wb.worksheets:
        meta[ws.title] = dict(merged=sorted(str(r) for r in ws.merged_cells.ranges),
                              cf=sorted(str(r.sqref) for r in ws.conditional_formatting),
                              filter=ws.auto_filter.ref, freeze=ws.freeze_panes,
                              dv=len(ws.data_validations.dataValidation))
        if ws.title in skip:
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is not None and not isinstance(v, (str, int, float, bool, dt.datetime, dt.date)):
                    v = str(getattr(v, "text", v))
                f = c.font
                sty = (c.number_format, bool(f.b), bool(f.i), str(f.color.rgb) if f.color is not None else None,
                       str(c.fill.fgColor.rgb) if c.fill.fill_type else None, c.alignment.horizontal)
                if v is None and sty == ("General", False, False, None, None, None):
                    continue
                cells[(ws.title, c.coordinate)] = (v,) + sty
    return meta, cells


def wb_diff(a, b, skip=()):
    ma, ca = wb_snapshot(a, skip)
    mb, cb = wb_snapshot(b, skip)
    out = [("meta", k, ma.get(k), mb.get(k)) for k in sorted(set(ma) | set(mb), key=str) if ma.get(k) != mb.get(k)]
    for k in sorted(set(ca) | set(cb)):
        if ca.get(k) != cb.get(k):
            out.append((k[0], k[1], ca.get(k), cb.get(k)))
    return out


def screener_rows(path):
    """Pharma rows by ticker, after checking the columns are still where bscreener writes (layout v2)."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    if "Pharma" not in wb.sheetnames:
        raise Stop("Screener has no Pharma tab")
    ws = wb["Pharma"]
    rows = {str(ws.cell(r, 2).value).strip(): r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value}
    n = len(rows)
    chk = {"F (price) holds numbers": sum(isinstance(ws.cell(r, 6).value, (int, float)) for r in rows.values()),
           "G = F x Main C4": sum(bool(re.match(rf"=F{r}\*\[\d+\]Main!\$C\$4$", str(ws.cell(r, 7).value)))
                                  for r in rows.values()),
           "J reads Main C11": sum("Main!$C$11" in str(ws.cell(r, 10).value) for r in rows.values()),
           "K = J / F - 1": sum(f"J{r}/F{r}" in str(ws.cell(r, 11).value) for r in rows.values())}
    bad = [f"{k} in {v} of {n} rows" for k, v in chk.items() if v < 0.95 * n]
    if bad:
        raise Stop("The Screener's Pharma columns aren't where the refresh expects them (layout v2: F price, "
                   "G market cap, J NPV / Share, K upside, L model): " + "; ".join(bad) +
                   ". If you moved columns, bscreener.py and refresh_daily.py need updating first.")
    return rows


def step_screener(ctx):
    if is_locked(SCREENER):
        ctx["state"]["screener"] = "locked"
        say("SKIPPED: Screener.xlsx is open in Excel (~$Screener.xlsx exists). Close it, then run "
            "refresh_daily.py --from screener" + (f" --date {ctx['iso']}" if ctx["args"].date else ""))
        return
    px = load_closes(ctx) if os.path.exists(ctx["closes_file"]) else {}
    rows = screener_rows(SCREENER)
    upd = {t: r for t, r in rows.items() if t in px}
    say(f"Pharma tab: {len(rows)} rows, {len(upd)} prices to set, F1 -> 'Close {ctx['iso']}'"
        + (f"; no close for {sorted(set(rows) - set(upd))[:10]}" if px and len(upd) < len(rows) else ""))
    if ctx["dry"]:
        say("  would back up Screener.xlsx, set F prices + F1, run bscreener.py and fixlinks_all.py, check that only "
            "column F changed, then install")
        return
    os.makedirs(f"{TOOLS}/screener/backups", exist_ok=True)
    stamp = dt.datetime.now().strftime("%H%M%S")
    backup = f"{TOOLS}/screener/backups/Screener_{ctx['ymd']}_{stamp}_before-prices.xlsx"
    shutil.copyfile(SCREENER, backup)
    say(f"  backup: {rel(backup)}")
    st = f"{ctx['run_dir']}/screener"
    os.makedirs(st, exist_ok=True)
    s1, s2, s3 = f"{st}/1_prices.xlsx", f"{st}/2_bscreener.xlsx", f"{st}/3_fixlinks.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(backup)
    ws = wb["Pharma"]
    for t, r in upd.items():
        ws.cell(r, 6).value = px[t]
    ws["F1"] = f"Close {ctx['iso']}"
    wb.save(s1)
    log = f"{ctx['run_dir']}/screener.log"
    run([PY, "bscreener.py", s1, s2], cwd=BASIC, log=log, echo=True)
    run([PY, "fixlinks_all.py", s2, s3], cwd=TOOLS, log=log, echo=True)
    allowed = {("Pharma", "F1")} | {("Pharma", f"F{r}") for r in upd.values()}
    diffs = wb_diff(backup, s3)
    bad = [d for d in diffs if d[:2] not in allowed]
    wrong = [d for d in diffs if d[:2] in allowed and d[1] != "F1" and not same((d[3] or (None,))[0],
                                                                                  px[ws.cell(int(d[1][1:]), 2).value.strip()])]
    if bad or wrong:
        for d in (bad + wrong)[:15]:
            say(f"  unexpected change: {d}")
        raise Stop(f"Screener check failed: {len(bad)} changes outside column F, {len(wrong)} wrong prices. "
                   f"Not installed; staged files in {rel(st)}/", 2)
    say(f"  check passed: {len(diffs)} cells changed, all in column F (prices) or F1")
    if is_locked(SCREENER) or md5(SCREENER) != md5(backup):
        raise Stop("Screener.xlsx was opened or changed while the refresh ran; not installed. Run --from screener.")
    copy_atomic(s3, SCREENER)
    ctx["state"]["screener"] = "updated"
    say("  installed Screener.xlsx")


# ---------------------------------------------------------------- 8 recs
def build_universe(ctx):
    px = load_closes(ctx) if os.path.exists(ctx["closes_file"]) else {}
    prices = {x["t"]: x["price"] for x in jload(TICKERS)}
    prices.update(px)
    U = []
    for t in todo_list():
        L, C = jload(f"{BASIC}/out/{t}_layout.json"), jload(f"{BASIC}/out/{t}_calc.json")
        kind = {"DCF": "Reviewed DCF" if L["P"].get("review") else "Basic DCF", "Net cash": "Net cash",
                "No data": "No data", "rNPV": "Pipeline rNPV"}[L["kind"]]
        nd = L["kind"] == "No data"
        r = dict(t=t, kind=kind, value=None if nd else C["main"]["C11"], price=L["price"],
                 up=None if nd else C["main"]["C12"], mcap=L["mcap"] or 0.0, flag=bool(L["flag"]))
        if L["kind"] == "DCF":
            r["m0"] = L["P"]["m0"]
        U.append(r)
    for t, kind in BUILT.items():
        v = main_vals(f"{TOOLS}/reprice_out/{t}.xlsx", ("C4", "C11", "C12"))
        U.append(dict(t=t, kind=kind, value=v["C11"], price=prices[t], up=v["C12"],
                      mcap=prices[t] * v["C4"] if isinstance(v["C4"], float) else None, flag=False))
    return U


def step_recs(ctx):
    rec_out = f"{TOOLS}/rec_{ctx['ymd']}.json"
    if ctx["dry"]:
        say(f"would rebuild universe.json, run rec.py universe.json {rel(rec_out)}, then rec_sheet.py + fixlinks on "
            "the Screener (only if the Screener step installed it), checking only the Recommendations sheet changed")
        return
    U = build_universe(ctx)
    shutil.copyfile(UNIVERSE, f"{ctx['run_dir']}/universe.before.json")
    jsave(UNIVERSE, U, indent=0)
    say(f"universe.json: {len(U)} rows")
    log = f"{ctx['run_dir']}/recs.log"
    out = run([PY, "rec.py", UNIVERSE, rec_out], cwd=TOOLS, log=log)
    say("  rec.py:", out.strip().splitlines()[0])
    if is_locked(SCREENER) or ctx["state"].get("screener") == "locked":
        ctx["state"]["recs"] = "skipped"
        say("SKIPPED the sheet: Screener.xlsx is open in Excel. Close it, then run refresh_daily.py --from screener")
        return
    import openpyxl
    f1 = openpyxl.load_workbook(SCREENER, read_only=True)["Pharma"]["F1"].value
    if f1 != f"Close {ctx['iso']}":
        ctx["state"]["recs"] = "skipped"
        raise Stop(f"The Screener's Pharma F1 is {f1!r}, not 'Close {ctx['iso']}': its prices aren't from this close, "
                   "so the Recommendations sheet wasn't rewritten. Run --from screener.")
    stamp = dt.datetime.now().strftime("%H%M%S")
    backup = f"{TOOLS}/screener/backups/Screener_{ctx['ymd']}_{stamp}_before-recs.xlsx"
    os.makedirs(os.path.dirname(backup), exist_ok=True)
    shutil.copyfile(SCREENER, backup)
    st = f"{ctx['run_dir']}/screener"
    os.makedirs(st, exist_ok=True)
    r1, r2 = f"{st}/4_rec_sheet.xlsx", f"{st}/5_fixlinks.xlsx"
    run([PY, "rec_sheet.py", backup, r1, rec_out, ctx["iso"]], cwd=TOOLS, log=log, echo=True)
    run([PY, "fixlinks_all.py", r1, r2], cwd=TOOLS, log=log, echo=True)
    diffs = wb_diff(backup, r2, skip={"Recommendations"})
    bad = [d for d in diffs if not (d[0] == "meta" and d[1] == "Recommendations")]
    a, b = (openpyxl.load_workbook(p, read_only=True) for p in (backup, r2))
    b3 = [x["Recommendations"]["B3"].value is not None if "Recommendations" in x.sheetnames else None for x in (a, b)]
    if bad or b3[0] != b3[1]:
        for d in bad[:15]:
            say(f"  unexpected change: {d}")
        raise Stop(f"Recommendations check failed: {len(bad)} changes outside the Recommendations sheet"
                   + (", B3 line changed" if b3[0] != b3[1] else "") + f". Not installed; staged in {rel(st)}/")
    say(f"  check passed: only the Recommendations sheet changed (tab {b.sheetnames.index('Recommendations') + 1} of "
        f"{len(b.sheetnames)}, B3 line {'kept' if b3[1] else 'still deleted'})")
    if is_locked(SCREENER) or md5(SCREENER) != md5(backup):
        raise Stop("Screener.xlsx was opened or changed while the refresh ran; not installed. Run --from recs.")
    copy_atomic(r2, SCREENER)
    ctx["state"]["recs"] = "updated"
    say("  installed Screener.xlsx")


# ---------------------------------------------------------------- 9 trade
def step_trade(ctx):
    cmd = [PY, "paper_trader.py", "run", "--prices", ctx["closes_file"], "--sheet", SCREENER] + \
        (["--force"] if ctx["args"].force_trade else [])
    if ctx["args"].no_trade:
        say("skipped (--no-trade)")
        return
    if ctx["dry"]:
        say("would run:", short(cmd))
        return
    import openpyxl
    f1 = openpyxl.load_workbook(SCREENER, read_only=True)["Pharma"]["F1"].value
    if f1 != f"Close {ctx['iso']}" or ctx["state"].get("recs") != "updated":
        say(f"SKIPPED: the Recommendations sheet isn't from the {ctx['iso']} close yet (Pharma F1 {f1!r}, "
            f"recs {ctx['state'].get('recs', 'not run')}). After closing Excel: refresh_daily.py --from screener")
        ctx["state"]["trade"] = "skipped"
        return
    env = dict(os.environ, TRADER_DIR=TRADER)
    led = jload(f"{TRADER}/ledger.json", {"positions": {}})
    applied = jload(SPLITS, {})
    for t, rec in ctx["state"].get("splits", {}).items():
        k = f"{t}:{rec['date']}"
        if t in led.get("positions", {}) and not applied.get(k, {}).get("trader"):
            run([PY, "paper_trader.py", "split", t, f"{rec['den'] / rec['num']:.10g}"], cwd=TRADER, echo=True, env=env)
            applied.setdefault(k, rec)["trader"] = True
            jsave(SPLITS, applied, indent=1, sort_keys=True)
    run(cmd, cwd=TRADER, log=f"{ctx['run_dir']}/trade.log", echo=True, env=env)
    ctx["state"]["trade"] = "done"


FUNCS = {"closes": step_closes, "splits": step_splits, "prices": step_prices, "validate": step_validate,
         "notes": step_notes, "install": step_install, "screener": step_screener, "recs": step_recs,
         "trade": step_trade}


# ---------------------------------------------------------------- main
def acquire_lock():
    p = f"{TOOLS}/.refresh.lock"
    for _ in range(2):
        try:
            fd = os.open(p, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, str(os.getpid()).encode())
            os.close(fd)
            return p
        except FileExistsError:
            try:
                pid = int(open(p).read().strip() or 0)
                os.kill(pid, 0)
                raise Stop(f"Another refresh is running (pid {pid}).")
            except (ValueError, ProcessLookupError):
                os.remove(p)                   # stale lock from a crashed run
    raise Stop("Couldn't take the refresh lock")


def main():
    global LOG
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="close date YYYY-MM-DD (default: today, after 4:00 pm ET)")
    ap.add_argument("--dry-run", action="store_true", help="print every step; write nothing")
    ap.add_argument("--accept-changes", action="store_true", help="install even if NPV/share values changed")
    ap.add_argument("--from", dest="start", choices=STEPS, default=STEPS[0], help="first step to run")
    ap.add_argument("--to", dest="end", choices=STEPS, default=STEPS[-1], help="last step to run")
    ap.add_argument("--refetch", action="store_true", help="fetch all closes again even if the file exists")
    ap.add_argument("--no-trade", action="store_true", help="skip the paper trader")
    ap.add_argument("--force-trade", action="store_true", help="pass --force to the paper trader")
    ap.add_argument("--jobs", type=int, default=6, help="parallel builds (default 6)")
    a = ap.parse_args()
    now = dt.datetime.now(NY)
    try:
        if a.date:
            day = dt.date.fromisoformat(a.date)
        else:
            day = now.date()
            if day.weekday() >= 5:
                raise Stop(f"{day} is a weekend. Pass --date for a past close.")
            if now.time() < dt.time(16, 0):
                raise Stop(f"It's {now:%H:%M} ET and the market closes at 4:00 pm. Pass --date for a past close.")
        if day > now.date():
            raise Stop(f"--date {day} is in the future")
    except Stop as e:
        print("STOPPED:", e)
        return e.code
    live = day == now.date() and now.time() >= dt.time(16, 0)
    ymd = day.strftime("%Y%m%d")
    ctx = dict(args=a, dry=a.dry_run, date=day, iso=day.isoformat(), ymd=ymd, live=live,
               stamp=f"{day:%b} {day.day}, {day.year}", closes_file=f"{TOOLS}/closes_{ymd}.json",
               run_dir=f"{TOOLS}/runs/{ymd}", state={}, manifest=None)
    i0, i1 = STEPS.index(a.start), STEPS.index(a.end)
    lock = None
    code = 0
    try:
        for p, what in ((PY, "venv python"), (PHARMA, "Pharma folder"), (SCREENER, "Screener.xlsx"),
                        (COMPANIES, "vault Companies folder"), (TRADER, "trader folder")):
            if not os.path.exists(p):
                raise Stop(f"Missing {what}: {p}")
        if not a.dry_run:
            lock = acquire_lock()
            os.makedirs(ctx["run_dir"], exist_ok=True)
            LOG = open(f"{ctx['run_dir']}/refresh.log", "a")
            sp = f"{ctx['run_dir']}/state.json"
            ctx["state"] = jload(sp, {})
        say(f"=== refresh_daily {'DRY RUN ' if a.dry_run else ''}close {ctx['iso']} | started {now:%Y-%m-%d %H:%M} ET | "
            f"steps {i0 + 1}-{i1 + 1} | {'quote feed (just closed)' if live else 'historical feed'}")
        say(f"    Stock Models: {ROOT}\n    Vault: {VAULT}")
        locks = sorted(glob.glob(f"{PHARMA}/~$*.xlsx")) + ([lock_of(SCREENER)] if is_locked(SCREENER) else [])
        if locks:
            say(f"    Excel has open: {[os.path.basename(p)[2:] for p in locks]} (those files won't be overwritten)")
        if i0 <= STEPS.index("prices"):
            ensure_manifest(ctx)
        for i in range(i0, i1 + 1):
            s = STEPS[i]
            say(f"\n[{i + 1}/9] {TITLES[s]}")
            FUNCS[s](ctx)
            if not a.dry_run:
                jsave(f"{ctx['run_dir']}/state.json", ctx["state"], indent=1, default=str)
    except Stop as e:
        say(f"\nSTOPPED: {e}")
        code = e.code
    except KeyboardInterrupt:
        say("\nInterrupted")
        code = 130
    finally:
        if lock and os.path.exists(lock):
            os.remove(lock)
    if not code:
        st = ctx["state"]
        notes = []
        if st.get("screener") == "locked":
            notes.append("Screener not updated (open in Excel)")
        if st.get("trade") == "skipped":
            notes.append("paper trader not run")
        if st.get("install_skipped"):
            notes.append(f"{len(st['install_skipped'])} files not installed (see step 6)")
        say(f"\n=== {'dry run finished: nothing was written' if a.dry_run else 'done'}"
            + (" | " + "; ".join(notes) if notes else ""))
    if LOG:
        LOG.write("\n")
        LOG.close()
    return code


if __name__ == "__main__":
    sys.exit(main())
