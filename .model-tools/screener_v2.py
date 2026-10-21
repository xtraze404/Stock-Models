"""Screener layout v2 (2026-10-20): drop no-data tickers (and their links), drop Shares/Cash/Debt (Market Cap and EV read
the model directly), keep Event, add Note(s), and write the Recommendations sheet.
usage: screener_v2.py SRC OUT"""
import copy, datetime as dt, json, os, re, sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList

D = os.path.dirname(os.path.abspath(__file__))
SRC, OUT = sys.argv[1], sys.argv[2]
U = {r["t"]: r for r in json.load(open(f"{D}/universe.json"))}
EARN = json.load(open(f"{D}/earnings_20261020.json"))
PD = json.load(open(f"{D}/pdufa_events.json"))
REC = json.load(open(f"{D}/rec_20261020.json"))
CLOSE = "2026-10-20"
NODATA = {t for t, r in U.items() if r["kind"] == "No data"}
EXTRA = {"QNCX": "Now IRulya Therapeutics (IRLA) after an Oct 9 merger; model shares and cash predate it",
         "REVB": "1-for-7 reverse split Oct 20", "MBRX": "1-for-20 reverse split Oct 20",
         "ANAB": None}

wb = openpyxl.load_workbook(SRC)
ws = wb["Pharma"]
FONT = ws["B3"].font
HFONT = ws["B2"].font
HBORDER = copy.copy(ws["B2"].border)

# ---------------------------------------------------------------- collect rows (current order)
rows = []
for r in range(3, ws.max_row + 1):
    t = ws.cell(r, 2).value
    if not t:
        continue
    t = str(t).strip()
    m = re.match(r"=\[(\d+)\]Main!\$C\$4", str(ws.cell(r, 9).value))
    assert m, (t, ws.cell(r, 9).value)
    rows.append(dict(t=t, vals=[ws.cell(r, c).value for c in range(2, 7)], n=int(m.group(1)),
                     M=ws.cell(r, 13).value, O=ws.cell(r, 15).value, P=ws.cell(r, 16).value,
                     style={c: copy.copy(ws.cell(r, c)._style) for c in range(2, 17)}))
keep = [x for x in rows if x["t"] not in NODATA]
dropped = [x["t"] for x in rows if x["t"] in NODATA]
old_links = wb._external_links
used = sorted({x["n"] for x in keep})
newidx = {n: i + 1 for i, n in enumerate(used)}
for x in keep:
    assert old_links[x["n"] - 1].file_link.Target.split("/")[-1] == f"{x['t']}.xlsx", x["t"]
wb._external_links = [old_links[n - 1] for n in used]
relink = lambda f, n: re.sub(rf"\[{n}\]", f"[{newidx[n]}]", f) if isinstance(f, str) else f

# ---------------------------------------------------------------- clear and rewrite the table
hdr_style = {c: copy.copy(ws.cell(2, c)._style) for c in range(2, 17)}
last_old = ws.max_row
for r in range(2, last_old + 1):
    for c in range(2, 17):
        cell = ws.cell(r, c)
        cell.value = None
        cell._style = copy.copy(ws.cell(1, 1)._style)
# header
heads = {2: "Ticker", 3: "Name", 4: "Industry", 5: "Exchange", 6: "Price", 7: "Market Cap", 8: "Enterprise V.",
         9: None, 10: "NPV / Share", 11: "Upside", 12: "Model", 13: "Event", 14: "Note(s)"}
src_col = {2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: 7, 8: 8, 9: 12, 10: 13, 11: 14, 12: 15, 13: 16, 14: 16}
for c, h in heads.items():
    cell = ws.cell(2, c)
    cell._style = copy.copy(hdr_style[src_col[c]])
    cell.value = h
ws.cell(2, 14).alignment = Alignment(horizontal="left")


def fmt_day(iso):
    d = dt.date.fromisoformat(iso)
    return f"{d:%b} {d.day}"


notes_n = events_n = 0
for i, x in enumerate(keep):
    r = 3 + i
    k = newidx[x["n"]]
    for j, v in enumerate(x["vals"]):
        ws.cell(r, 2 + j).value = v
    for c in range(2, 15):
        ws.cell(r, c)._style = copy.copy(x["style"][src_col[c]])
    ws.cell(r, 7).value = f"=F{r}*[{k}]Main!$C$4"
    ws.cell(r, 8).value = f"=G{r}-[{k}]Main!$C$7+[{k}]Main!$C$8"
    ws.cell(r, 10).value = relink(x["M"], x["n"])
    ws.cell(r, 11).value = f'=IFERROR(J{r}/F{r}-1,"")'
    ws.cell(r, 12).value = relink(x["O"], x["n"])
    # Event: keep what is there (PDUFA dates); otherwise the next earnings date
    e = EARN.get(x["t"])
    ev = x["P"]
    if ev is None and isinstance(e, dict) and e["date"] >= CLOSE:
        ev = dt.datetime.fromisoformat(e["date"])
        events_n += 1
    ws.cell(r, 13).value = ev
    ws.cell(r, 13).number_format = "yyyy-mm-dd"
    ws.cell(r, 13).alignment = Alignment(horizontal="center")
    parts = []
    if x["t"] in PD and PD[x["t"]]["date"]:
        drug = PD[x["t"]]["drug"]
        if x["t"] == "ANAB":
            drug = "Quimilza (Vanda; ANAB royalty)"
        parts.append(f"PDUFA {PD[x['t']]['date'].rsplit(',', 1)[0]}: {drug}")
    if isinstance(e, dict) and e["date"] >= CLOSE:
        s = f"Earnings {fmt_day(e['date'])}"
        if e["when"]:
            s += f" ({e['when']})"
        if e["status"] == "estimated":
            s += " (est.)"
        parts.append(s)
    if EXTRA.get(x["t"]):
        parts.append(EXTRA[x["t"]])
    nc = ws.cell(r, 14)
    nc.value = " · ".join(parts) if parts else None
    nc.font = copy.copy(FONT)
    nc.alignment = Alignment(horizontal="left")
    notes_n += bool(parts)
last = 2 + len(keep)
for c in (15, 16):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 9
widths = {"I": 3, "J": 12, "K": 9, "L": 25, "M": 12, "N": 58}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
ws.conditional_formatting = ConditionalFormattingList()
ws.conditional_formatting.add(f"K3:K{last}", FormulaRule(formula=["AND(ISNUMBER(K3),K3<0)"], font=Font(color="9C0006"),
                              fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"), stopIfTrue=True))
ws.conditional_formatting.add(f"K3:K{last}", FormulaRule(formula=["AND(ISNUMBER(K3),K3>0)"], font=Font(color="006100"),
                              fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"), stopIfTrue=True))
ws.auto_filter.ref = f"B2:N{last}"
ws.freeze_panes = "C3"

# ---------------------------------------------------------------- Recommendations sheet
if "Recommendations" in wb.sheetnames:
    del wb["Recommendations"]
R = wb.create_sheet("Recommendations", index=wb.sheetnames.index("Pharma") + 1)
F = lambda bold=False, italic=False, color="000000", size=12: Font(name=FONT.name, size=size, bold=bold, italic=italic, color=color)
GREEN, RED, MUTED = "006100", "9C0006", "595959"
put = lambda a, v, fmt=None, **kw: (setattr(R[a], "value", v), setattr(R[a], "font", F(**kw)),
                                    fmt and setattr(R[a], "number_format", fmt))
R.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGHIJKLMNOPQR", (5, 8, 34, 7, 9, 14, 12, 10, 9, 12, 11, 10, 9, 7, 11, 12, 60)):
    R.column_dimensions[col].width = w
put("B2", "Recommendations: model-weighted long/short portfolio (100% gross)", bold=True, size=14)
put("B3", f"Weights from the Stock Models valuations at the {CLOSE} close. Model output, not investment advice; "
          "I'm not a licensed advisor.", italic=True, color=MUTED)
P = REC["positions"]
L = [p for p in P if p["side"] == "Long"]
S = [p for p in P if p["side"] == "Short"]
wl = sum(p["weight"] for p in L)
ws_ = -sum(p["weight"] for p in S)
summ = [("Long book", wl, "0.0%"), ("Short book", -ws_, "0.0%;-0.0%"), ("Gross exposure", wl + ws_, "0.0%"),
        ("Net exposure", wl - ws_, "0.0%"), ("Long positions", len(L), "0"), ("Short positions", len(S), "0"),
        ("Weighted upside, longs", sum(p["weight"] * p["up"] for p in L) / wl, "+0%;-0%"),
        ("Weighted upside, shorts", sum(-p["weight"] * p["up"] for p in S) / ws_, "+0%;-0%"),
        ("Valuation models in universe", REC["n_universe"], "0"),
        ("Positions after filters", len(P), "0")]
put("B5", "Summary", bold=True)
for i, (lab, v, fmt) in enumerate(summ):
    put(f"C{6 + i}", lab)
    R.merge_cells(f"C{6 + i}:E{6 + i}")
    put(f"F{6 + i}", v, fmt, bold=i < 4)
c = REC["cfg"]
rules = ["How the weights are set",
         "Universe: the 168 models with a valuation (DCF or rNPV). Net-cash and no-data tickers are left out.",
         f"Liquidity: price at least ${c['min_price']:.0f} and market cap at least ${c['min_mcap'] / 1000:.0f}B, so micro-caps and penny stocks can't drive the book.",
         f"Signal: long if upside is +{c['thresh']:.0%} or more, short if -{c['thresh']:.0%} or less; upside clipped to +100% / -60%.",
         "Score = upside × model confidence × size factor. Confidence: reviewed or full DCF 1.0, rNPV 0.7, basic DCF 0.5, basic DCF with an operating loss today 0.25.",
         "Size factor rises from 0.25 at $1B to 1.0 at $100B+ (log scale).",
         f"Books: {c['long_book']:.0%} long / {c['short_book']:.0%} short. Caps: 8% for $50B+, 6% for $10–50B, 4% for $3–10B, 2% for $1–3B; positions under 0.5% are dropped.",
         "Basic DCFs use GAAP margins and no pipeline, so they lean negative for acquisitive big pharma; that's why they get half weight."]
for i, s in enumerate(rules):
    put(f"H{5 + i}", s, bold=i == 0, color="000000" if i == 0 else MUTED)
h0 = 17
hd = ["#", "Ticker", "Company", "Side", "Weight", "Model", "Value / share", "Price", "Upside", "Mkt cap ($B)",
      "Confidence", "Size factor", "Cap", "Score", "Live upside", "Event", "Note(s)"]
thin = Border(bottom=Side(style="thin"))
for i, h in enumerate(hd):
    cell = R.cell(h0, 2 + i)
    cell.value = h
    cell.font = F(bold=True)
    cell.border = thin
    cell.alignment = Alignment(horizontal="left" if h in ("Ticker", "Company", "Model", "Note(s)") else "center")
names = {x["t"]: x["vals"][1] for x in keep}
r = h0 + 1
for grp, label in ((L, "Longs"), (S, "Shorts")):
    R.cell(r, 3).value = label
    R.cell(r, 3).font = F(bold=True)
    r += 1
    for j, p in enumerate(grp, start=1):
        vals = [j, p["t"], names.get(p["t"], ""), p["side"], p["weight"], p["kind"], p["value"], p["price"], p["up"],
                p["mcap"] / 1000, p["cf"], p["sf"], p["cap"], p["score"]]
        fmts = ["0", None, None, None, "0.0%;-0.0%", None, "$#,##0.00", "$#,##0.00", "+0%;-0%", "#,##0.0", "0.00",
                "0.00", "0%", "0.000"]
        for i, (v, fm) in enumerate(zip(vals, fmts)):
            cell = R.cell(r, 2 + i)
            cell.value = v
            cell.font = F(color=(GREEN if p["side"] == "Long" else RED) if i in (3, 4) else "000000", bold=i == 4)
            if fm:
                cell.number_format = fm
            cell.alignment = Alignment(horizontal="left" if i in (1, 2, 5) else "center")
        mt = f'MATCH($C{r},Pharma!$B$3:$B${last},0)'
        R.cell(r, 16).value = f'=IFERROR(INDEX(Pharma!$K$3:$K${last},{mt}),"")'
        R.cell(r, 16).number_format = "+0%;-0%"
        R.cell(r, 17).value = f'=IFERROR(IF(INDEX(Pharma!$M$3:$M${last},{mt})="","",INDEX(Pharma!$M$3:$M${last},{mt})),"")'
        R.cell(r, 17).number_format = "yyyy-mm-dd"
        R.cell(r, 18).value = f'=IFERROR(T(INDEX(Pharma!$N$3:$N${last},{mt})),"")'
        for cc in (16, 17, 18):
            R.cell(r, cc).font = F()
            R.cell(r, cc).alignment = Alignment(horizontal="left" if cc == 18 else "center")
        r += 1
    R.cell(r, 4).value = f"Total {label.lower()}"
    R.cell(r, 4).font = F(bold=True)
    R.cell(r, 6).value = sum(p["weight"] for p in grp)
    R.cell(r, 6).number_format = "0.0%;-0.0%"
    R.cell(r, 6).font = F(bold=True, color=GREEN if label == "Longs" else RED)
    R.cell(r, 6).border = Border(top=Side(style="thin"))
    R.cell(r, 6).alignment = Alignment(horizontal="center")
    r += 2
R.freeze_panes = f"B{h0 + 1}"
R.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f"kept {len(keep)} rows, dropped {len(dropped)}: {dropped}")
print(f"links {len(old_links)} -> {len(wb._external_links)} | events filled {events_n} | notes {notes_n} | "
      f"recommendations {len(L)} long / {len(S)} short | table to row {last}")
