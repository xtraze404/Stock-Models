"""Paper-trade the Screener's Recommendations sheet in a local simulated account.

Schwab's Trader API only reaches live accounts, so this keeps its own ledger (start: $100,000, like thinkorswim
paperMoney), rebalances it to the sheet's target weights at each run, and writes the day's orders in thinkorswim terms
(BUY / SELL / SELL SHORT / BUY TO COVER) so they can be mirrored in paperMoney by hand.

usage:
  paper_trader.py run [--dry-run] [--force] [--prices FILE.json] [--sheet Screener.xlsx]
  paper_trader.py status
  paper_trader.py split TICKER NEW_PER_OLD      e.g. "split REVB 0.142857" after a 1-for-7 reverse split
  paper_trader.py reset [--cash 100000]

Files (in this folder): config.json, ledger.json, orders/<date>.csv, positions.csv, nav.csv.
Model output, not investment advice."""
import argparse, csv, datetime as dt, json, math, os, re, sys, time, urllib.request

HERE = os.environ.get("TRADER_DIR") or os.path.dirname(os.path.abspath(__file__))
DEFAULT_SHEET = "/Users/eric/Documents/Stock Models/Screener.xlsx"
DEFAULTS = {
    "start_cash": 100000.0,
    "band": 0.0075,             # resize a position only when it is this far (share of NAV) from target
    "min_trade_usd": 250.0,     # ignore smaller resizes (exits and new names always trade)
    "slippage_bps": 5.0,        # fills: buys at price x (1 + bps), sells at price x (1 - bps)
    "borrow_rate": 0.02,        # annual short borrow fee on short market value, accrued per calendar day
    "borrow_overrides": {},     # e.g. {"SMMT": 0.08} for hard-to-borrow names
    "max_sheet_age_days": 4,    # refuse to trade on a stale sheet (calendar days) unless --force
    "split_guard": 0.6,         # a held name moving more than this since the last run is held back (possible split)
    "max_missing_prices": 0.1,  # abort if more than this share of tickers has no price
    "aliases": {"QNCX": "IRLA"},
}
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 "
                    "Safari/537.36", "Accept": "application/json"}


# ---------------------------------------------------------------- config / ledger
def load_cfg():
    p = os.path.join(HERE, "config.json")
    cfg = dict(DEFAULTS)
    if os.path.exists(p):
        cfg.update(json.load(open(p)))
    else:
        json.dump(DEFAULTS, open(p, "w"), indent=2)
    return cfg


def load_ledger(cfg):
    p = os.path.join(HERE, "ledger.json")
    if os.path.exists(p):
        return json.load(open(p))
    return {"cash": cfg["start_cash"], "positions": {}, "last_prices": {}, "last_run": None, "trades": [],
            "fees_paid": 0.0, "created": dt.date.today().isoformat()}


def save_ledger(L):
    p = os.path.join(HERE, "ledger.json")
    tmp = p + ".tmp"
    json.dump(L, open(tmp, "w"), indent=1)
    os.replace(tmp, p)


# ---------------------------------------------------------------- targets from the sheet
def read_targets(sheet):
    """{ticker: signed weight} from the Recommendations sheet, plus its as-of date."""
    import openpyxl
    wb = openpyxl.load_workbook(sheet, data_only=True)
    if "Recommendations" not in wb.sheetnames:
        raise SystemExit("No Recommendations sheet in " + sheet)
    R = wb["Recommendations"]
    asof = None
    for v in [R.cell(r, c).value for r in range(1, 7) for c in range(1, 8)] + \
             ([wb["Pharma"]["F1"].value] if "Pharma" in wb.sheetnames else []):
        m = re.search(r"(\d{4}-\d{2}-\d{2})", str(v or ""))
        if m:                                  # sheet text, else the Pharma price header "Close YYYY-MM-DD"
            asof = dt.date.fromisoformat(m.group(1))
            break
    hdr = next((r for r in range(1, 60) if R.cell(r, 3).value == "Ticker" and R.cell(r, 6).value == "Weight"), None)
    if hdr is None:
        raise SystemExit("Recommendations sheet: header row (Ticker / Weight) not found")
    tg = {}
    for r in range(hdr + 1, R.max_row + 1):
        t, side, w = R.cell(r, 3).value, R.cell(r, 5).value, R.cell(r, 6).value
        if isinstance(t, str) and t.strip() and side in ("Long", "Short") and isinstance(w, (int, float)):
            w = abs(float(w)) * (1 if side == "Long" else -1)
            if w:
                tg[t.strip().upper()] = w
    return tg, asof


# ---------------------------------------------------------------- prices
def quote(t, aliases):
    """Last price: Nasdaq quote feed (regular-session close after 4pm ET), Cboe as fallback. -> (price, source)"""
    sym = aliases.get(t, t)
    for cls in ("stocks", "etf"):
        for attempt in range(3):
            try:
                d = json.load(urllib.request.urlopen(urllib.request.Request(
                    f"https://api.nasdaq.com/api/quote/{sym}/info?assetclass={cls}", headers=UA), timeout=20)).get("data")
                if not d:
                    break
                sec, pri = d.get("secondaryData") or {}, d.get("primaryData") or {}
                if "Closed at" in (sec.get("lastTradeTimestamp") or "") and sec.get("lastSalePrice"):
                    return float(sec["lastSalePrice"].replace("$", "").replace(",", "")), "nasdaq close"
                if pri.get("lastSalePrice"):
                    tag = "nasdaq close" if d.get("marketStatus") == "Closed" else "nasdaq " + (d.get("marketStatus") or "")
                    return float(pri["lastSalePrice"].replace("$", "").replace(",", "")), tag.strip()
                break
            except Exception:
                time.sleep(1.5 * (attempt + 1))
    try:
        d = json.load(urllib.request.urlopen(urllib.request.Request(
            f"https://cdn.cboe.com/api/global/delayed_quotes/options/{sym}.json", headers=UA), timeout=20))["data"]
        p = d.get("current_price") or d.get("close")
        if p:
            return float(p), "cboe"
    except Exception:
        pass
    return None, None


def get_prices(tickers, cfg, prices_file=None):
    if prices_file:
        raw = json.load(open(prices_file))
        out = {}
        for t in tickers:
            v = raw.get(t)
            v = v.get("px") if isinstance(v, dict) else v
            out[t] = (float(v), "file") if v else (None, None)
        return out
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(3) as ex:
        return dict(zip(tickers, ex.map(lambda t: quote(t, cfg["aliases"]), tickers)))


# ---------------------------------------------------------------- core
def nav_of(L, px):
    return L["cash"] + sum(q * px[t] for t, q in L["positions"].items() if px.get(t))


def plan_orders(L, targets, px, cfg, held_back):
    """Whole-share orders that move the ledger to the targets. -> list of dicts"""
    nav = nav_of(L, px)
    orders = []
    names = sorted(set(targets) | set(L["positions"]))
    for t in names:
        p = px.get(t)
        cur = L["positions"].get(t, 0)
        if not p or t in held_back:
            continue
        w = targets.get(t, 0.0)
        tgt = int(w * nav / p) if w else 0          # truncate toward zero: never overshoot the target
        delta = tgt - cur
        if delta == 0:
            continue
        exiting, opening = tgt == 0, cur == 0
        drift = abs(cur * p - tgt * p) / nav
        if not (exiting or opening or (cur > 0) != (tgt > 0)):
            if drift < cfg["band"] or abs(delta) * p < cfg["min_trade_usd"]:
                continue
        reason = "exit" if exiting else "new" if opening else "flip" if (cur > 0) != (tgt > 0) else "rebalance"
        legs = []
        if cur > 0 and tgt < 0:
            legs = [("SELL", cur), ("SELL SHORT", -tgt)]
        elif cur < 0 and tgt > 0:
            legs = [("BUY TO COVER", -cur), ("BUY", tgt)]
        elif tgt >= 0 and cur >= 0:
            legs = [("BUY", delta)] if delta > 0 else [("SELL", -delta)]
        else:
            legs = [("SELL SHORT", -delta)] if delta < 0 else [("BUY TO COVER", delta)]
        for action, q in legs:
            orders.append({"action": action, "ticker": t, "qty": q, "price": p, "reason": reason,
                           "target_weight": w, "target_shares": tgt})
    # sells and covers first, then buys and new shorts (mirrors how you'd free up buying power)
    order = {"SELL": 0, "BUY TO COVER": 1, "SELL SHORT": 2, "BUY": 3}
    orders.sort(key=lambda o: (order[o["action"]], o["ticker"]))
    return orders, nav


def apply_orders(L, orders, cfg, day):
    s = cfg["slippage_bps"] / 1e4
    for o in orders:
        buy = o["action"] in ("BUY", "BUY TO COVER")
        fill = o["price"] * (1 + s if buy else 1 - s)
        val = fill * o["qty"]
        L["cash"] += -val if buy else val
        L["positions"][o["ticker"]] = L["positions"].get(o["ticker"], 0) + (o["qty"] if buy else -o["qty"])
        if L["positions"][o["ticker"]] == 0:
            del L["positions"][o["ticker"]]
        o["fill"], o["value"] = round(fill, 4), round(val, 2)
        L["trades"].append(dict(o, date=day))


def accrue_borrow(L, px, cfg, day):
    if not L["last_run"]:
        return 0.0
    days = (dt.date.fromisoformat(day) - dt.date.fromisoformat(L["last_run"])).days
    if days <= 0:
        return 0.0
    fee = 0.0
    for t, q in L["positions"].items():
        if q < 0:
            p = px.get(t) or L["last_prices"].get(t)
            if p:
                fee += -q * p * cfg["borrow_overrides"].get(t, cfg["borrow_rate"]) * days / 360
    L["cash"] -= fee
    L["fees_paid"] = L.get("fees_paid", 0.0) + fee
    return fee


def write_outputs(L, orders, px, day, nav_before, fee):
    os.makedirs(os.path.join(HERE, "orders"), exist_ok=True)
    with open(os.path.join(HERE, "orders", f"{day}.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Action", "Symbol", "Quantity", "Est. price", "Est. value", "Reason", "Target weight", "Target shares"])
        for o in orders:
            w.writerow([o["action"], o["ticker"], o["qty"], f"{o['price']:.2f}", f"{o['qty'] * o['price']:.2f}", o["reason"],
                        f"{o['target_weight']:.2%}", o["target_shares"]])
    nav = nav_of(L, px)
    with open(os.path.join(HERE, "positions.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Symbol", "Shares", "Side", "Price", "Market value", "Weight of NAV"])
        for t, q in sorted(L["positions"].items(), key=lambda kv: -abs(kv[1] * (px.get(kv[0]) or 0))):
            p = px.get(t) or L["last_prices"].get(t, 0)
            w.writerow([t, q, "Long" if q > 0 else "Short", f"{p:.2f}", f"{q * p:.2f}", f"{q * p / nav:.2%}"])
    lm = sum(q * px[t] for t, q in L["positions"].items() if q > 0 and px.get(t))
    sm = -sum(q * px[t] for t, q in L["positions"].items() if q < 0 and px.get(t))
    p = os.path.join(HERE, "nav.csv")
    new = not os.path.exists(p)
    with open(p, "a", newline="") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["Date", "NAV", "Cash", "Long MV", "Short MV", "Gross %", "Net %", "Return since last run",
                        "Borrow fee", "Orders"])
        ret = nav / nav_before - 1 if nav_before else 0
        w.writerow([day, f"{nav:.2f}", f"{L['cash']:.2f}", f"{lm:.2f}", f"{sm:.2f}", f"{(lm + sm) / nav:.1%}",
                    f"{(lm - sm) / nav:.1%}", f"{ret:.2%}", f"{fee:.2f}", len(orders)])
    return nav, lm, sm


def cmd_run(a, cfg):
    day = dt.date.today().isoformat()
    L = load_ledger(cfg)
    if L["last_run"] == day and not (a.force or a.dry_run):
        raise SystemExit(f"Already ran today ({day}). Use --force to run again.")
    targets, asof = read_targets(a.sheet)
    if asof and (dt.date.today() - asof).days > cfg["max_sheet_age_days"] and not a.force:
        raise SystemExit(f"Recommendations sheet is as of {asof}, older than {cfg['max_sheet_age_days']} days. "
                         "Refresh it first, or use --force.")
    gross = sum(abs(w) for w in targets.values())
    if not 0.9 <= gross <= 1.1:
        print(f"Warning: target weights add up to {gross:.1%} gross, not 100%.")
    names = sorted(set(targets) | set(L["positions"]))
    q = get_prices(names, cfg, a.prices)
    px = {t: v[0] for t, v in q.items() if v[0]}
    missing = [t for t in names if t not in px]
    if len(missing) > cfg["max_missing_prices"] * max(1, len(names)):
        raise SystemExit(f"Too many tickers without a price ({len(missing)}): {missing[:10]}")
    # held positions without a price keep their last price for NAV and are not traded
    for t in missing:
        if t in L["positions"] and t in L["last_prices"]:
            px[t] = L["last_prices"][t]
    held_back = set(missing)
    for t, qty in L["positions"].items():
        lp = L["last_prices"].get(t)
        if lp and t in q and q[t][0] and abs(q[t][0] / lp - 1) > cfg["split_guard"]:
            held_back.add(t)
            px[t] = lp      # value it at the last price until the split is recorded, so NAV and sizing stay right
            print(f"Held back {t}: price moved {q[t][0] / lp - 1:+.0%} since the last run. If it split, run "
                  f"'split {t} <new shares per old share>' and run again.")
    nav_before = nav_of(L, L["last_prices"]) if L["last_run"] else L["cash"]   # value at the end of the last run
    W = dict(L)
    W["last_prices"] = dict(L["last_prices"])
    fee = accrue_borrow(W, px, cfg, day) if not a.dry_run else 0.0
    orders, nav = plan_orders(W, targets, px, cfg, held_back)
    print(f"Sheet as of {asof} | {len(targets)} targets | NAV ${nav:,.2f} | {len(orders)} orders"
          + (f" | borrow fee ${fee:,.2f}" if fee else "") + (" | DRY RUN" if a.dry_run else ""))
    for o in orders:
        print(f"  {o['action']:12s} {o['qty']:>7,d} {o['ticker']:6s} @ ~${o['price']:,.2f}  ({o['reason']}, target "
              f"{o['target_weight']:+.2%})")
    if missing:
        print(f"No price for: {missing} (not traded)")
    if a.dry_run:
        return
    apply_orders(W, orders, cfg, day)
    W["last_prices"].update(px)
    W["last_run"] = day
    nav_after, lm, sm = write_outputs(W, orders, px, day, nav_before, fee)
    save_ledger(W)
    print(f"Done. NAV ${nav_after:,.2f} | long ${lm:,.0f} / short ${sm:,.0f} | cash ${W['cash']:,.0f} | "
          f"orders saved to orders/{day}.csv")


def cmd_status(cfg):
    L = load_ledger(cfg)
    px = L["last_prices"]
    nav = nav_of(L, px)
    print(f"Last run {L['last_run']} | NAV ${nav:,.2f} (at last-run prices) | cash ${L['cash']:,.2f} | "
          f"{len(L['positions'])} positions | borrow fees paid ${L.get('fees_paid', 0):,.2f}")
    for t, q in sorted(L["positions"].items(), key=lambda kv: -abs(kv[1] * px.get(kv[0], 0))):
        print(f"  {t:6s} {q:>8,d}  {q * px.get(t, 0) / nav:+.2%}")


def cmd_split(a, cfg):
    L = load_ledger(cfg)
    t, r = a.ticker.upper(), float(a.ratio)
    if t not in L["positions"]:
        raise SystemExit(f"No position in {t}")
    old = L["positions"][t]
    new = int(old * r) if old > 0 else -int(-old * r)
    if t in L["last_prices"]:
        L["last_prices"][t] = L["last_prices"][t] / r
    L["positions"][t] = new
    L["trades"].append({"date": dt.date.today().isoformat(), "action": "SPLIT", "ticker": t, "qty": new, "ratio": r})
    save_ledger(L)
    print(f"{t}: {old} -> {new} shares (cash-in-lieu of fractions ignored)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--dry-run", action="store_true")
    r.add_argument("--force", action="store_true")
    r.add_argument("--prices")
    r.add_argument("--sheet", default=DEFAULT_SHEET)
    sub.add_parser("status")
    s = sub.add_parser("split")
    s.add_argument("ticker")
    s.add_argument("ratio")
    z = sub.add_parser("reset")
    z.add_argument("--cash", type=float)
    a = ap.parse_args()
    cfg = load_cfg()
    if a.cmd == "run":
        cmd_run(a, cfg)
    elif a.cmd == "status":
        cmd_status(cfg)
    elif a.cmd == "split":
        cmd_split(a, cfg)
    elif a.cmd == "reset":
        p = os.path.join(HERE, "ledger.json")
        if os.path.exists(p):
            os.replace(p, p.replace(".json", f".{dt.datetime.now():%Y%m%d-%H%M%S}.bak.json"))
        if a.cash:
            cfg["start_cash"] = a.cash
            json.dump(cfg, open(os.path.join(HERE, "config.json"), "w"), indent=2)
        print("Ledger reset (old one kept as a .bak.json).")


if __name__ == "__main__":
    main()
