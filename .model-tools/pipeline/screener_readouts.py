"""Add pipeline catalysts to the Screener's Note(s) column (N) for the pipeline-rNPV companies: the next trial readout
(ClinicalTrials.gov primary completion date) and any pending PDUFA date. Existing note text is kept; earlier
"Readout" / "PDUFA" entries written by this script are replaced.  usage: screener_readouts.py SRC OUT"""
import json, os, re, sys
import openpyxl

D = os.path.dirname(os.path.abspath(__file__))
B = os.path.join(os.path.dirname(D), "basic")
STG = {4: "approved", 3.5: "filed", 3.0: "Phase 3", 2.5: "Phase 2/3", 2.0: "Phase 2", 1.5: "Phase 1/2", 1.0: "Phase 1"}
SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb["Pharma"]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
NC = hdr["Note(s)"]
n = 0
for r in range(3, ws.max_row + 1):
    t = ws.cell(r, 2).value
    if not t:
        continue
    t = str(t).strip()
    lp = f"{B}/out/{t}_layout.json"
    old = ws.cell(r, NC).value or ""
    parts = [p for p in old.split(" · ") if p and not re.match(r"(Readout ~|PDUFA \d{4}-|Pipeline:)", p)]
    if os.path.exists(lp):
        L = json.load(open(lp))
        if L.get("kind") == "rNPV":
            PR = L["programs"]
            nxt = sorted(((p["next_readout"], p) for p in PR if p.get("next_readout") and p["stage"] < 4), key=lambda x: x[0])
            if nxt:
                d_, p = nxt[0]
                parts.append(f"Readout ~{d_}: {p['bucket']}, {STG.get(p['stage'], p['stage'])}")
            for p in PR:
                m = re.search(r"PDUFA (\w{3}) (\d{1,2}), (20\d\d)", p.get("reg", ""))
                if m and p["stage"] == 3.5 and not any(x.startswith("PDUFA ") for x in parts):
                    parts.append(f"PDUFA {m.group(3)}-{['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].index(m.group(1)) + 1:02d}"
                                 f"-{int(m.group(2)):02d}: {p['bucket']}")
    new = " · ".join(parts) or None
    if new != (old or None):
        ws.cell(r, NC).value = new
        n += 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("notes updated", n)
